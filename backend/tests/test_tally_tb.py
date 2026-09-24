from decimal import Decimal
from pathlib import Path

import openpyxl
import pytest

from app.tally_tb import GROUP, LEDGER, SUB_GROUP, find_tb_sheet, format_inr, parse_trial_balance

VENUS = Path(__file__).resolve().parents[2] / "samples" / "Venus_Agencies_2025-26.xlsx"
GRAND_TOTAL = Decimal("101252474.34")  # ₹10,12,52,474.34 on both sides

# Every italic ledger in the Venus TB: (row, name, group path, Dr, Cr)
VENUS_LEDGERS = [
    (13, "Advance Income Tax", "Capital Account", "195000", "0"),
    (14, "Capital Gain on Sale on Gold Bond", "Capital Account", "0", "65850"),
    (15, "GODS A/C", "Capital Account", "0", "3384"),
    (16, "Interest on FD", "Capital Account", "0", "78607"),
    (17, "T L Venkatesh", "Capital Account", "0", "12051800.97"),
    (19, "T.L.Venkatesh (H.U.F )", "Loans (Liability)", "0", "1669000"),
    (22, "G S T", "Current Liabilities > Duties & Taxes", "0", "207.83"),
    (23, "R.M.C .60%", "Current Liabilities > Duties & Taxes", "1714.68", "0"),
    (26, "Car Accunnt", "Fixed Assets", "127263", "0"),
    (27, "Computer  A/C", "Fixed Assets", "4086", "0"),
    (28, "Furnitures & Fixtures", "Fixed Assets", "541", "0"),
    (29, "Sanghavi Opera Flat (50% Share)", "Fixed Assets", "1614925", "0"),
    (30, "Scooter", "Fixed Assets", "66580", "0"),
    (31, "Shop Building", "Fixed Assets", "1300555", "0"),
    (32, "Weighing Scale A/C", "Fixed Assets", "740", "0"),
    (34, "G.M.C.Bank Shares A/C", "Investments", "45911", "0"),
    (35, "NAMASTE CO-OPERATIVE BANK (SHARES)", "Investments", "10250", "0"),
    (36, "N.S.C", "Investments", "694163", "0"),
    (40, "Membership", "Current Assets > Deposits (Asset)", "13508", "0"),
    (41, "Mysore Silk Cloth Merchant Co Op .,", "Current Assets > Deposits (Asset)", "300211", "0"),
    (42, "R.M.C.Bank Gurantee A/c", "Current Assets > Deposits (Asset)", "9832", "0"),
    (43, "Tax Saving Fixed Deposite Account", "Current Assets > Deposits (Asset)", "154797", "0"),
    (48, "Bank of Baroda", "Current Assets > Bank Accounts", "1029943.56", "0"),
    (49, "Union Bank of India", "Current Assets > Bank Accounts", "174594.64", "0"),
    (50, "Inventory Difference Account", "Current Assets", "623000", "0"),
    (52, "Gst Sales @ 5%", "Sales Accounts", "0", "13990194.28"),
    (53, "Gst Sales Nil Rated", "Sales Accounts", "0", "68484428.83"),
    (54, "IGST NIL RATED", "Sales Accounts", "0", "1644000"),
    (55, "I G S T Sales @ 5%", "Sales Accounts", "0", "12430"),
    (57, "Discount", "Purchase Accounts", "0", "589476.89"),
    (58, "Gst Purchases @18%", "Purchase Accounts", "78000", "0"),
    (59, "Gst Purchases @ 5%", "Purchase Accounts", "10423494.57", "0"),
    (60, "Gst Purchases  Nil Rated", "Purchase Accounts", "27748501", "0"),
    (61, "I G S T Puchase @5%", "Purchase Accounts", "3052181", "0"),
    (62, "I G S T Purchases Nil Rated", "Purchase Accounts", "39889407.02", "0"),
    (64, "Cooli", "Direct Expenses", "190891", "0"),
    (65, "Fright Charges", "Direct Expenses", "1125257", "0"),
    (66, "Round Off", "Direct Expenses", "0", "0.94"),
    (70, "Audit Fees", "Indirect Expenses", "53600", "0"),
    (71, "Bank Chargers", "Indirect Expenses", "19041.61", "0"),
    (72, "CORPORATION TAX", "Indirect Expenses", "25423", "0"),
    (73, "Depreciation", "Indirect Expenses", "35425", "0"),
    (74, "Electricity  Charges", "Indirect Expenses", "2965", "0"),
    (75, "G S T Arrears", "Indirect Expenses", "1647", "0"),
    (76, "Membership Fee", "Indirect Expenses", "5000", "0"),
    (77, "Salary  A/c", "Indirect Expenses", "504000", "0"),
    (78, "SHOP EXPENCESS", "Indirect Expenses", "101055", "0"),
    (79, "Shop Insurance A/c", "Indirect Expenses", "3540", "0"),
    (80, "Tally Software Services", "Indirect Expenses", "6500", "0"),
    (81, "Telephone Charges", "Indirect Expenses", "31786.74", "0"),
    (82, "Vehicle Insurance", "Indirect Expenses", "6080", "0"),
    (83, "Vehicle Maintainance", "Indirect Expenses", "104140.38", "0"),
]

# Groups Tally shows collapsed: they hold a balance but have no ledgers under them.
VENUS_COLLAPSED = [
    (24, "Sundry Creditors", SUB_GROUP, "Current Liabilities", "55066", "2040093.6"),
    (38, "Opening Stock", SUB_GROUP, "Current Assets", "3664516.55", "0"),
    (44, "Loans & Advances (Asset)", SUB_GROUP, "Current Assets", "20000", "0"),
    (45, "Sundry Debtors", SUB_GROUP, "Current Assets", "6720081.63", "0"),
    (46, "Cash-in-hand", SUB_GROUP, "Current Assets", "1017259.96", "0"),
    (68, "Closing Stock", GROUP, "", "0", "3398447.26"),
    (84, "Profit & Loss A/c", GROUP, "", "0", "623000"),
]


@pytest.fixture(scope="module")
def tb():
    return parse_trial_balance(VENUS)


def test_finds_tb_sheet_by_name():
    wb = openpyxl.load_workbook(VENUS, data_only=True)
    ws, (name_col, debit_col, first_row) = find_tb_sheet(wb)
    assert ws.title == "TB"
    assert (name_col, debit_col, first_row) == (1, 2, 12)


def test_finds_tb_sheet_by_content_when_renamed():
    wb = openpyxl.load_workbook(VENUS, data_only=True)
    wb["TB"].title = "Ledger balances"
    tb = parse_trial_balance(wb)
    assert tb.sheet_name == "Ledger balances"
    assert tb.total_dr == GRAND_TOTAL


def test_company_details(tb):
    assert tb.company == "venus agencies 2025-26"
    assert tb.address == ["14/30 M.G.B.COMPLEX", "APMC YARD YPR", "BANGALORE-560022"]
    assert tb.gstin == "29ABGPV9302D1ZH"
    assert tb.period == "1-Apr-25 to 31-Mar-26"
    assert (tb.debit_column, tb.credit_column) == ("B", "C")


def test_every_ledger_read(tb):
    got = [(r.row, r.name, r.group_path, r.dr, r.cr) for r in tb.ledgers]
    expected = [(row, name, path, Decimal(dr), Decimal(cr)) for row, name, path, dr, cr in VENUS_LEDGERS]
    assert got == expected


def test_collapsed_groups_hold_balances(tb):
    by_row = {r.row: r for r in tb.rows}
    for row, name, kind, path, dr, cr in VENUS_COLLAPSED:
        r = by_row[row]
        assert (r.name, r.kind, r.group_path, r.dr, r.cr) == (name, kind, path, Decimal(dr), Decimal(cr))
        assert r.children == []


def test_row_kinds(tb):
    kinds = {r.name: r.kind for r in tb.rows}
    assert kinds["Current Assets"] == GROUP
    assert kinds["Bank Accounts"] == SUB_GROUP
    assert kinds["Duties & Taxes"] == SUB_GROUP
    assert kinds["Bank of Baroda"] == LEDGER
    assert kinds["Indirect Incomes"] == GROUP  # empty group, no balance
    assert len(tb.rows) == 73  # rows 12 to 84


def test_totals_equal_on_both_sides(tb):
    assert tb.total_dr == GRAND_TOTAL
    assert tb.total_cr == GRAND_TOTAL
    assert tb.is_balanced
    assert format_inr(tb.total_dr) == "10,12,52,474.34"


def test_balances_add_up_without_double_counting(tb):
    # Ledgers plus collapsed groups: every rupee in the TB counted exactly once.
    # Closing Stock is a memo line: Tally leaves it out of the Grand Total.
    assert len(tb.balances) == len(VENUS_LEDGERS) + len(VENUS_COLLAPSED)
    assert [r.name for r in tb.balances if r.memo] == ["Closing Stock"]
    counted = [r for r in tb.balances if not r.memo]
    assert sum(r.dr for r in counted) == GRAND_TOTAL
    assert sum(r.cr for r in counted) == GRAND_TOTAL


def test_no_problems_in_venus_tb(tb):
    assert tb.problems == []


def test_unbalanced_tb_is_reported():
    wb = openpyxl.load_workbook(VENUS, data_only=True)
    wb["TB"]["C85"] = 101252000  # tamper with the credit Grand Total
    tb = parse_trial_balance(wb)
    assert not tb.is_balanced
    assert any("does not equal credit total" in p and "474.34" in p for p in tb.problems)


def test_group_subtotal_mismatch_is_reported():
    wb = openpyxl.load_workbook(VENUS, data_only=True)
    wb["TB"]["B48"] = 1029000  # Bank of Baroda no longer adds up to Bank Accounts
    tb = parse_trial_balance(wb)
    assert any("'Bank Accounts'" in p for p in tb.problems)


def test_workbook_without_tb_raises():
    wb = openpyxl.Workbook()
    wb.active["A1"] = "Nothing here"
    with pytest.raises(ValueError, match="No Tally trial balance"):
        parse_trial_balance(wb)


@pytest.mark.parametrize("amount, text", [
    (Decimal("101252474.34"), "10,12,52,474.34"),
    (Decimal("13674167.85"), "1,36,74,167.85"),
    (Decimal("623000"), "6,23,000.00"),
    (Decimal("207.83"), "207.83"),
    (Decimal("-3384"), "-3,384.00"),
])
def test_format_inr(amount, text):
    assert format_inr(amount) == text
