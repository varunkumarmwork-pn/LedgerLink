"""BUILD MODE on the Venus TB: default formats, only the ticked sheets, live formulas to the TB."""
import json

import openpyxl
import pytest

from app.build import SHEETS, assign_lines, build_workbook
from app.learn import learn_workbook
from app.review import build_review
from app.tally_tb import parse_trial_balance
from app.templates import line_for
from tb_samples import VENUS, make_venus_tb_with_suspense

# Venus FY 2025-26 from the TB alone (default mapping, so not the CA's hand-made statements):
# revenue 8,41,31,053.11 less purchases, direct expenses, change in stock (opening 36,64,516.55
# less closing 33,98,447.26), salary, bank charges, depreciation, other expenses and tax.
PROFIT_BEFORE_TAX = 1071949.33
PROFIT = 1046526.33


def build(tmp_path, entity, sheets=SHEETS, tb=VENUS, name="built.xlsx"):
    out = tmp_path / name
    facts = build_workbook(tb, entity, "Venus Agencies", sheets, out)
    result = learn_workbook(out)
    mapping = {**json.loads(json.dumps(result.mapping())), "build": facts}
    report = json.loads(json.dumps(result.report()))
    review = build_review(out, mapping, mapping, report)
    return {"out": out, "facts": facts, "result": result, "review": review, "wb": openpyxl.load_workbook(out)}


def figure(built, sheet, label):
    return next(round(c.value, 2) for k, c in built["result"].cells.items()
                if k.strip("'").startswith(sheet) and c.label == label and c.financial_year == "2025-26")


# ---------------------------------------------------------------- default mapping from Tally groups

@pytest.mark.parametrize("names, entity, line", [
    (["Bank of Baroda", "Bank Accounts", "Current Assets"], "Proprietorship", "cash_bank"),
    (["Sundry Debtors", "Current Assets"], "Proprietorship", "trade_receivables"),
    (["G S T", "Duties & Taxes", "Current Liabilities"], "Proprietorship", "other_current_liabilities"),
    (["T L Venkatesh", "Capital Account"], "Proprietorship", "capital"),
    (["Share capital", "Capital Account"], "Private limited company", "share_capital"),
    (["Salary  A/c", "Indirect Expenses"], "Proprietorship", "employee"),
    (["Bank Chargers", "Indirect Expenses"], "Proprietorship", "finance"),
    (["SHOP EXPENCESS", "Indirect Expenses"], "Proprietorship", "other_expenses"),
    (["Fright Charges", "Direct Expenses"], "Proprietorship", "direct_expenses"),
    (["Fright Charges", "Direct Expenses"], "Private limited company", "other_expenses"),
    (["Unknown receipt", "Suspense A/c"], "Proprietorship", None),
])
def test_default_mapping(names, entity, line):
    assert line_for(names, entity) == line


def test_every_venus_ledger_has_a_line():
    lines, unmapped = assign_lines(parse_trial_balance(VENUS), "Proprietorship")
    assert unmapped == []
    assert [r.name for r in lines["inventory_change"]] == ["Opening Stock", "Closing Stock"]


# ---------------------------------------------------------------- non-corporate, all sheets

@pytest.fixture(scope="module")
def proprietor(tmp_path_factory):
    return build(tmp_path_factory.mktemp("prop"), "Proprietorship")


def test_all_ticked_sheets_in_order(proprietor):
    assert proprietor["wb"].sheetnames == ["Balance Sheet", "Profit & Loss", "BS Schedules", "P&L Notes", "PPE", "TB",
                                           "Accounting policies"]
    assert proprietor["facts"]["template"] == "Non-corporate (ICAI Guidance Note)"


def test_live_formulas_to_the_tb(proprietor):
    wb = proprietor["wb"]
    assert wb["Balance Sheet"]["B7"].value == "Proprietor's Capital Account"
    assert wb["Balance Sheet"]["D7"].value == "='BS Schedules'!C12"  # line -> note total
    notes = {wb["BS Schedules"].cell(r, 2).value: wb["BS Schedules"].cell(r, 3).value for r in range(1, 80)}
    assert notes["Bank of Baroda"] == "=TB!B48-TB!C48"  # note row -> TB ledger
    assert notes["Add: Profit for the year"] == "='Profit & Loss'!D18"  # profit -> capital
    assert wb["Balance Sheet"]["D17"].value == "=PPE!F11"  # PPE closing WDV


def test_figures_and_checks(proprietor):
    assert figure(proprietor, "Profit & Loss", "Profit before Tax") == PROFIT_BEFORE_TAX
    assert figure(proprietor, "Profit & Loss", "Profit for the Year") == PROFIT
    assert figure(proprietor, "Balance Sheet", "Cash and Bank Balances") == 2221798.16
    checks = {c["id"]: c for c in proprietor["review"]["checks"]}
    assert checks["tb"]["ok"] and checks["bs"]["ok"] and checks["unused"]["ok"]
    assert proprietor["review"]["items"] == []  # every ledger mapped, nothing to review


# ---------------------------------------------------------------- company, Schedule III

def test_company_uses_schedule_iii(tmp_path):
    built = build(tmp_path, "Private limited company")
    ws = built["wb"]["Balance Sheet"]
    labels = [ws.cell(r, 2).value for r in range(1, 40) if ws.cell(r, 2).value]
    assert "I. EQUITY AND LIABILITIES" in labels and "Share Capital" in labels
    assert "Total Equity and Liabilities" in labels
    assert built["facts"]["template"] == "Company (Schedule III, Division I)"
    assert figure(built, "Profit & Loss", "Other Expenses") == 1632461.18  # includes direct expenses
    assert figure(built, "Profit & Loss", "Profit for the Year") == PROFIT
    assert {c["id"]: c["ok"] for c in built["review"]["checks"]}["bs"]


# ---------------------------------------------------------------- only the ticked sheets

def test_only_ticked_sheets_link_straight_to_the_tb(tmp_path):
    built = build(tmp_path, "Partnership firm", sheets=["Balance Sheet", "Profit & Loss"])
    wb = built["wb"]
    assert wb.sheetnames == ["Balance Sheet", "Profit & Loss", "TB"]
    ws = wb["Balance Sheet"]
    row = next(r for r in range(1, 40) if ws.cell(r, 2).value == "Trade Receivables")
    assert ws.cell(row, 4).value == "=(TB!B45-TB!C45)"  # no notes sheet: straight to the TB
    assert ws["B7"].value == "Partners' Capital Accounts"
    assert figure(built, "Profit & Loss", "Profit for the Year") == PROFIT
    assert {c["id"]: c["ok"] for c in built["review"]["checks"]}["bs"]
    assert built["review"]["items"] == []


def test_no_sheets_ticked_is_refused(tmp_path):
    with pytest.raises(ValueError, match="Tick at least one sheet"):
        build_workbook(VENUS, "Proprietorship", "Venus", [], tmp_path / "x.xlsx")


# ---------------------------------------------------------------- unmapped ledgers are red

def test_unmapped_ledger_shows_red_in_review(tmp_path):
    tb = make_venus_tb_with_suspense(tmp_path / "tb.xlsx")
    built = build(tmp_path, "Proprietorship", tb=tb)
    assert built["facts"]["unmapped"] == ["Suspense A/c > Unknown receipt"]
    item = next(i for i in built["review"]["items"] if i["id"] == "unused:Suspense A/c > Unknown receipt")
    assert item["tone"] == "red"
    assert "Its Tally group has no line in the default format." in item["text"]
    assert built["review"]["overrides"]["TB!A86"]["status"] == "r"
    checks = {c["id"]: c for c in built["review"]["checks"]}
    assert checks["bs"]["label"] == "Balance sheet difference ₹1,000.00"  # the unmapped credit


# ---------------------------------------------------------------- through the API

def test_build_through_the_api(client, storage):
    body = {"name": "Venus Agencies", "entity_type": "Private limited company", "financial_year": "2025-26"}
    client_id = client.post("/api/clients", json=body).json()["id"]
    sheets = ["Balance Sheet", "Profit & Loss", "BS Schedules"]
    started = client.post(f"/api/clients/{client_id}/build", files={"file": (VENUS.name, VENUS.read_bytes())},
                          data={"sheets": sheets})
    assert started.status_code == 202
    job = client.get(f"/api/jobs/{started.json()['job_id']}").json()
    assert job["status"] == "done", job["error"]
    assert job["action"] == "Building" and job["titles"][1] == "Put ledgers on statement lines"

    workbook = client.get(f"/api/clients/{client_id}/workbook").json()
    assert (workbook["source"], workbook["status"], workbook["financial_year"]) == ("build", "draft", "2025-26")
    names = [workbook["snapshot"]["sheets"][s]["name"] for s in workbook["snapshot"]["sheetOrder"]]
    assert names == sheets + ["TB"]
    assert {c["id"]: c["ok"] for c in workbook["checks"]}["bs"] and workbook["review"] == []
    assert client.get("/api/clients").json()[0]["review_status"] == "Ready to finalise"
    assert client.get(f"/api/clients/{client_id}/workbook/download").status_code == 200
    assert workbook["next_financial_year"] == "2026-27"
    early = client.post(f"/api/clients/{client_id}/rollforward", files={"file": (VENUS.name, VENUS.read_bytes())})
    assert early.status_code == 409 and early.json()["detail"].startswith("Finalise FY 2025-26 first")


def test_build_rejects_unknown_sheets(client, storage):
    body = {"name": "X", "entity_type": "LLP", "financial_year": "2025-26"}
    client_id = client.post("/api/clients", json=body).json()["id"]
    response = client.post(f"/api/clients/{client_id}/build", files={"file": (VENUS.name, VENUS.read_bytes())},
                           data={"sheets": ["Cash Flow"]})
    assert response.status_code == 422
