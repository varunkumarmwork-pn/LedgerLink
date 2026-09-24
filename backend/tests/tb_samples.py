"""Builds a FY 2026-27 Tally TB for Venus to test roll-forward (CLAUDE.md, Phase 5).

Changes from the FY 2025-26 TB:
  * new ledger "Interest on SB Account" (Cr 5,101.88) near the top of Capital Account,
    so almost every row below moves down one
  * new ledger "HDFC Bank" (Dr 50,000) under Bank Accounts; Cash-in-hand is 50,000 lower
  * "Bank Chargers" renamed to "Bank Charges"
  * "Tally Software Services" (Dr 6,500) no longer in the TB
  * Opening Stock = last year's closing stock 40,21,447.26
  * Sundry Creditors (Cr) absorbs the difference so the TB still balances
  * sheet named "TB 26-27", period 1-Apr-26 to 31-Mar-27

Run as a script to write samples/Venus_TB_2026-27_test.xlsx for trying it in the browser.
"""
from copy import copy
from decimal import Decimal
from pathlib import Path

import openpyxl

from app.tally_tb import GROUP, LEDGER, SUB_GROUP, parse_trial_balance

VENUS = Path(__file__).resolve().parents[2] / "samples" / "Venus_Agencies_2025-26.xlsx"
SAMPLE_OUT = VENUS.parent / "Venus_TB_2026-27_test.xlsx"
LAST_CLOSING_STOCK = Decimal("4021447.26")


class Row:
    def __init__(self, name, kind, indent, dr, cr, children=None, memo=False):
        self.name, self.kind, self.indent, self.dr, self.cr = name, kind, indent, dr, cr
        self.children = children or []
        self.memo = memo


def _tree(rows):
    return [Row(r.name, r.kind, r.indent, r.dr, r.cr, _tree(r.children), r.memo) for r in rows]


def _find(rows, name):
    for r in rows:
        if r.name == name:
            return r, rows
        found = _find(r.children, name)
        if found:
            return found
    return None


def _total(row):
    """Groups carry the sum of their rows, as in Tally."""
    if row.children:
        for child in row.children:
            _total(child)
        row.dr = sum((c.dr for c in row.children), Decimal("0.00"))
        row.cr = sum((c.cr for c in row.children), Decimal("0.00"))


def _flatten(rows):
    for r in rows:
        yield r
        yield from _flatten(r.children)


def make_modified_venus_tb(out_path: Path) -> Path:
    tb = parse_trial_balance(VENUS)
    top = _tree([r for r in tb.rows if not r.path])

    gods, capital = _find(top, "GODS A/C")
    capital.insert(capital.index(gods) + 1, Row("Interest on SB Account", LEDGER, 2, Decimal("0.00"), Decimal("5101.88")))
    union, banks = _find(top, "Union Bank of India")
    banks.insert(banks.index(union) + 1, Row("HDFC Bank", LEDGER, 3, Decimal("50000.00"), Decimal("0.00")))
    _find(top, "Cash-in-hand")[0].dr -= Decimal("50000.00")
    _find(top, "Bank Chargers")[0].name = "Bank Charges"
    tally, expenses = _find(top, "Tally Software Services")
    expenses.remove(tally)
    _find(top, "Opening Stock")[0].dr = LAST_CLOSING_STOCK

    for row in top:
        _total(row)
    total_dr = sum((r.dr for r in top if not r.memo), Decimal("0.00"))
    total_cr = sum((r.cr for r in top if not r.memo), Decimal("0.00"))
    creditors = _find(top, "Sundry Creditors")[0]
    creditors.cr += total_dr - total_cr  # keep Dr = Cr
    for row in top:
        _total(row)
    total = sum((r.dr for r in top if not r.memo), Decimal("0.00"))
    assert total == sum((r.cr for r in top if not r.memo), Decimal("0.00"))

    _write(top, total, out_path)
    return out_path


def _write(top, total, out_path):
    source = openpyxl.load_workbook(VENUS)["TB"]
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "TB 26-27"
    for row in source.iter_rows(max_row=11):  # company name, address, period, headings
        for c in row:
            target = ws.cell(c.row, c.column, c.value)
            _copy_style(c, target)
    ws["A1"] = "venus agencies 2026-27"
    ws["A7"] = ws["B9"] = "1-Apr-26 to 31-Mar-27"
    ws["B8"] = "venus agencies 2026-27"
    for letter, dim in source.column_dimensions.items():
        ws.column_dimensions[letter].width = dim.width

    style_rows = {GROUP: 12, SUB_GROUP: 21, LEDGER: 13}  # rows in the source with each kind's formatting
    r = 12
    for row in _flatten(top):
        template = style_rows[row.kind]
        for col, value in ((1, row.name), (2, row.dr or None), (3, row.cr or None)):
            cell = ws.cell(r, col, float(value) if isinstance(value, Decimal) else value)
            _copy_style(source.cell(template, col), cell)
        ws.cell(r, 1).alignment = copy(ws.cell(r, 1).alignment).__class__(indent=row.indent)
        r += 1
    ws.cell(r, 1, "Grand Total")
    ws.cell(r, 2, float(total))
    ws.cell(r, 3, float(total))
    for col in (1, 2, 3):
        _copy_style(source.cell(85, col), ws.cell(r, col))
    wb.save(out_path)


def _copy_style(src, dst):
    dst.font, dst.border, dst.fill = copy(src.font), copy(src.border), copy(src.fill)
    dst.alignment, dst.number_format = copy(src.alignment), src.number_format


def make_venus_tb_with_suspense(out_path: Path) -> Path:
    """Venus FY 2025-26 TB plus a "Suspense A/c" group (Cr 1,000), which has no default line
    in BUILD MODE. Cash-in-hand is 1,000 higher so the TB still balances."""
    wb = openpyxl.load_workbook(VENUS)
    ws = wb["TB"]
    ws.insert_rows(85, 2)  # before Grand Total
    for row, name, bold, indent in ((85, "Suspense A/c", True, 0), (86, "Unknown receipt", False, 2)):
        ws.cell(row, 1, name).font = copy(ws["A12" if bold else "A13"].font)
        ws.cell(row, 1).alignment = copy(ws["A13"].alignment).__class__(indent=indent)
        ws.cell(row, 3, 1000.0).font = copy(ws["C12" if bold else "C14"].font)
    for address in ("B46", "B37", "B87", "C87"):  # Cash-in-hand, Current Assets, Grand Total
        ws[address] = ws[address].value + 1000
    wb.save(out_path)
    return out_path


if __name__ == "__main__":
    print(make_modified_venus_tb(SAMPLE_OUT))
