"""ROLL-FORWARD MODE: build next year's workbook from last year's and a new Tally TB.

1. Start from last year's workbook as the CA left it (edits included).
2. This year's figures become the prior-year column as fixed numbers; headings and dates
   move on one year.
3. The TB sheet is replaced by the new TB, and every TB formula is rewritten to the row
   where its ledger now is, found by NAME:
     same     same group path and name                     -> green
     moved    same name, different group                   -> amber
     renamed  found by the AI layer (ai.py): the client's past
              choices, then name/meaning similarity          -> amber, with reason and confidence
     missing  not in the new TB (the formula uses 0)        -> red
4. Openings are carried forward: "As per last Balance Sheet" = last year's closing capital,
   PPE opening = last year's closing WDV (additions reset to 0).

Returns facts about what happened; review.py turns them into colours, review items and checks.
"""
import re
from copy import copy
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path

import openpyxl
from openpyxl.cell.cell import MergedCell
from openpyxl.formula.tokenizer import Token
from openpyxl.worksheet.worksheet import Worksheet

from app import ai
from app import formula as fx
from app.excel import effective_workbook, learn_effective
from app.formula import cell_key
from app.learn import MANUAL, TB_LINKED, LearnResult, Progress, _formula_text, _fy_start_from_period, _label
from app.tally_tb import TBRow, TrialBalance, format_inr, parse_trial_balance

SAME, MOVED, RENAMED, MISSING = "same", "moved", "renamed", "missing"
PLACEHOLDER = re.compile(r"\{TB:([^|}]+)\|(Dr|Cr)\}")
RANGE_PLACEHOLDER = re.compile(r"\{TB:([^|}]+)\|(Dr|Cr)\}:\{TB:([^|}]+)\|(Dr|Cr)\}")
OPENING_LABEL = re.compile(r"as per last balance sheet|opening balance|balance b/?f|brought forward", re.I)
MONTHS = "jan|feb|mar|apr|may|jun|jul|aug|sep|oct|nov|dec"


def roll_forward(base_path: str | Path, base_mapping: dict, new_tb_path: str | Path, out_path: str | Path,
                 progress: Progress | None = None, history: dict | None = None) -> dict:
    """Write next year's workbook to out_path. Returns the roll-forward facts (JSON-ready)."""
    progress = progress or Progress()
    base = learn_effective(base_path, base_mapping.get("edits", {}))
    new_tb = parse_trial_balance(new_tb_path)
    base_year = int(base.financial_year[:4])
    new_year = _fy_start_from_period(new_tb.period)
    if new_year is not None and new_year != base_year + 1:
        raise ValueError(f"This trial balance is for FY {_fy(new_year)}, but the next year is FY {_fy(base_year + 1)}.")
    progress.step_done(0, _tb_summary(new_tb))

    matches = match_ledgers(base.tb.rows, new_tb.rows, history)
    counts = {how: sum(1 for m in matches.values() if m["how"] == how) for how in (SAME, MOVED, RENAMED, MISSING)}
    new_only = [r.key for r in new_tb.rows if r.key not in {m["new"] for m in matches.values()}]
    progress.step_done(1, f"{counts[SAME]} same, {counts[RENAMED] + counts[MOVED]} renamed or moved, "
                          f"{counts[MISSING]} missing, {len(new_only)} new.")

    wb = effective_workbook(base_path, base_mapping.get("edits", {}))
    tb_sheet = base.tb.sheet_name
    facts = {
        "from_financial_year": base.financial_year,
        "financial_year": _fy(base_year + 1),
        "matches": matches,
        "new_ledgers": new_only,
        "templates": {},  # cell -> formula template with LAST year's ledger keys
        "manual": [],  # typed amounts kept from last year
        "carried": [],  # openings carried forward
        "ppe": [],  # PPE sheets whose additions were reset
        "suggestions": [],  # typed amounts that equal TB ledgers, with a formula to link them
        "last_year_issues": base.inconsistencies,
        "last_closing_stock": _closing_stock(base),
    }

    current_cols, prior_pairs = _year_columns(base, base_year)
    _move_current_to_prior(wb, base, prior_pairs)
    _carry_capital(wb, base, current_cols, facts)
    _carry_ppe(wb, base, facts, base_year + 1)
    _rewrite_tb_formulas(wb, base, matches, new_tb, tb_sheet, current_cols, facts, progress)
    _keep_manual(base, current_cols, facts)
    fresh = {c["cell"] for c in facts["carried"]} | {cell for p in facts["ppe"] for cell in p["reset"]}
    facts["suggestions"] = [x for x in link_suggestions(base.manual_cells, base.tb) if x["cell"] not in fresh]
    _replace_tb_sheet(wb[tb_sheet], openpyxl.load_workbook(new_tb_path)[new_tb.sheet_name])
    for ws in wb.worksheets:
        if ws.title != tb_sheet:
            _shift_years(ws, base_year)
    wb.save(out_path)
    progress.step_done(2, f"{len(facts['templates'])} TB formulas rewritten, {len(facts['carried'])} openings "
                          f"carried forward.")
    return facts


# ---------------------------------------------------------------- matching ledgers by name

def match_ledgers(old_rows: list[TBRow], new_rows: list[TBRow], history: dict | None = None) -> dict[str, dict]:
    """{old key: {"new": new key | None, "how": same|moved|renamed|missing, "score": 0..1, ...}}
    AI matches (renamed) also carry "confidence", "source" and "reason"; they stay amber until
    the CA approves them."""
    new_by_key = {r.key: r for r in new_rows}
    taken: set[str] = set()
    matches: dict[str, dict] = {}

    for old in old_rows:  # same group path and name
        if old.key in new_by_key:
            matches[old.key] = {"new": old.key, "how": SAME, "score": 1.0}
            taken.add(old.key)

    for old in old_rows:  # same name, another group
        if old.key in matches:
            continue
        same_name = [r for r in new_rows if r.key not in taken and _norm(r.name) == _norm(old.name)]
        if same_name:
            matches[old.key] = {"new": same_name[0].key, "how": MOVED, "score": 1.0}
            taken.add(same_name[0].key)

    for old in old_rows:  # the AI layer: past choices, then similarity (name, meaning, group, amount)
        if old.key in matches:
            continue
        candidates = [_ai_row(r) for r in new_rows if r.key not in taken]
        found = ai.match_ledger(_ai_row(old), candidates, history)
        if found:
            new_name = new_by_key[found["new"]].name
            matches[old.key] = {"new": found["new"], "how": RENAMED, "score": found["confidence"],
                                "confidence": found["confidence"], "source": found["source"],
                                "reason": ai.reason(old.name, f"now “{new_name}”", found)}
            taken.add(found["new"])
        else:
            matches[old.key] = {"new": None, "how": MISSING, "score": 0.0}
    return matches


def _ai_row(row: TBRow) -> dict:
    return {"key": row.key, "name": row.name, "group": row.group_path, "amount": float(row.dr or row.cr)}


def _norm(name: str) -> str:
    return re.sub(r"[^a-z0-9]+", " ", name.lower()).strip()


# ---------------------------------------------------------------- rewriting TB formulas

def render_template(template: str, tb: TrialBalance, tb_sheet: str, resolve) -> tuple[str, list[str]]:
    """Turn {TB:<ledger key>|Dr} placeholders into cell references in the new TB.

    resolve(old key) -> new key or None. A ledger that cannot be found becomes 0.
    Returns (formula, keys that were missing).
    """
    rows = {r.key: r for r in tb.rows}
    sheet = tb_sheet if tb_sheet.replace("_", "").isalnum() else "'" + tb_sheet.replace("'", "''") + "'"
    missing = []

    def address(old_key, side):
        row = rows.get(resolve(old_key) or "")
        if row is None:
            missing.append(old_key)
            return None
        return f"{tb.debit_column if side == 'Dr' else tb.credit_column}{row.row}"

    def one_range(m):
        start, end = address(m[1], m[2]), address(m[3], m[4])
        return f"{sheet}!{start}:{end}" if start and end else "0"

    def one(m):
        ref = address(m[1], m[2])
        return f"{sheet}!{ref}" if ref else "0"

    formula = PLACEHOLDER.sub(one, RANGE_PLACEHOLDER.sub(one_range, template))
    return formula, missing


def _rewrite_tb_formulas(wb, base: LearnResult, matches, new_tb, tb_sheet, current_cols, facts, progress):
    resolve = lambda old_key: matches.get(old_key, {}).get("new")  # noqa: E731
    for key, cell in base.cells.items():
        if cell.category != TB_LINKED or not cell.template:
            continue
        sheet, address = _split(key)
        formula, _ = render_template(cell.template, new_tb, tb_sheet, resolve)
        wb[sheet][address].value = formula
        facts["templates"][key] = cell.template
        names = ", ".join(dict.fromkeys(_name(new_tb, resolve(link.ledger)) or link.ledger.split(" > ")[-1]
                                        for link in cell.tb_links))
        progress.traced(key, formula, names)


def _name(tb: TrialBalance, key: str | None) -> str | None:
    return next((r.name for r in tb.rows if r.key == key), None)


# ---------------------------------------------------------------- prior year and headings

def _year_columns(base: LearnResult, year: int):
    """Current-year columns per sheet, and (current, prior) column pairs."""
    current, pairs = {}, {}
    for sheet, cols in base.year_columns.items():
        this = sorted(c for c, y in cols.items() if y == year)
        last = sorted(c for c, y in cols.items() if y == year - 1)
        current[sheet] = this
        pairs[sheet] = list(zip(this, last))
    return current, pairs


def _move_current_to_prior(wb, base: LearnResult, pairs):
    """This year's figures become next year's prior-year column, as fixed numbers."""
    for sheet, column_pairs in pairs.items():
        ws = wb[sheet]
        for current_col, prior_col in column_pairs:
            for row in range(1, ws.max_row + 1):
                now = base.cells.get(cell_key(sheet, row, current_col))
                prior = ws.cell(row, prior_col)
                if isinstance(prior, MergedCell):
                    continue
                if now is not None:
                    prior.value = round(now.value, 2)
                elif ws.cell(row, current_col).value is None and not isinstance(prior.value, str):
                    prior.value = None


def _shift_years(ws: Worksheet, year: int):
    """Headings and dates move on one year: "FY 2025-26" -> "FY 2026-27", "31.03.2026" -> "31.03.2027"."""
    for row in ws.iter_rows():
        for c in row:
            if isinstance(c.value, str) and not c.value.startswith("="):
                c.value = shift_text(c.value, year)
            elif isinstance(c.value, (datetime, date)) and c.value.year in (year, year + 1):
                c.value = _add_year(c.value)


def shift_text(text: str, year: int) -> str:
    """Move FY labels and dates that belong to the last two years on by one year."""
    def fy(m):
        start = int(m[1])
        if start in (year - 1, year) and int(m[3]) == (start + 1) % 100:
            return f"{start + 1}{m[2]}{(start + 2) % 100:02d}"
        return m[0]

    def year_only(m):
        y = int(m[2])
        return f"{m[1]}{y + 1}" if y in (year, year + 1) else m[0]

    text = re.sub(r"(?<!\d)(20\d{2})(\s*[-–]\s*)(\d{2})(?!\d)", fy, text)
    text = re.sub(r"((?<!\d)\d{1,2}([./-])\d{1,2}\2)(20\d{2})(?!\d)",
                  lambda m: m[1] + (str(int(m[3]) + 1) if int(m[3]) in (year, year + 1) else m[3]), text)
    text = re.sub(rf"((?:{MONTHS})[a-z]*\.?\s+(?:\d{{1,2}}(?:st|nd|rd|th)?,?\s+)?)(20\d{{2}})(?!\d)",
                  year_only, text, flags=re.I)
    text = re.sub(rf"(\d{{1,2}}(?:st|nd|rd|th)?\s+(?:{MONTHS})[a-z]*\.?,?\s+)(20\d{{2}})(?!\d)",
                  year_only, text, flags=re.I)
    return text


def _add_year(value):
    try:
        return value.replace(year=value.year + 1)
    except ValueError:  # 29 February
        return value.replace(year=value.year + 1, day=28)


# ---------------------------------------------------------------- carrying openings forward

def _carry_capital(wb, base: LearnResult, current_cols, facts):
    """ "As per last Balance Sheet" = the closing balance it led to last year."""
    for sheet, cols in current_cols.items():
        ws = wb[sheet]
        for col in cols:
            for row in range(1, ws.max_row + 1):
                opening = cell_key(sheet, row, col)
                label = _label(ws, row, col, has_year_heading=True)
                if not OPENING_LABEL.search(label or "") or opening not in base.cells:
                    continue
                closing = next((k for k in _cells_below(base, sheet, row, col) if opening in base.precedents.get(k, [])),
                               None)
                if closing is None:
                    continue
                amount = round(base.cells[closing].value, 2)
                ws.cell(row, col).value = amount
                facts["carried"].append({"cell": opening, "label": label, "from": closing, "amount": amount})


def _cells_below(base: LearnResult, sheet: str, row: int, col: int):
    ws_cells = [(k, c) for k, c in base.cells.items() if _split(k)[0] == sheet]
    below = []
    for k, _ in ws_cells:
        _, address = _split(k)
        r, c = _row_col(address)
        if c == col and r > row:
            below.append((r, k))
    return [k for _, k in sorted(below)]


def _carry_ppe(wb, base: LearnResult, facts, new_year: int):
    """PPE: opening = last year's closing; additions and deletions reset to 0."""
    for ws in wb.worksheets:
        header = _ppe_header(ws)
        if header is None:
            continue
        head_row, opening_col, closing_col, input_cols = header
        reset = []
        for row in range(head_row + 1, ws.max_row + 1):
            label = _label(ws, row, closing_col, has_year_heading=True).lower()
            if label.startswith("total"):
                break
            closing = base.cells.get(cell_key(ws.title, row, closing_col))
            if closing is None:
                continue
            opening = cell_key(ws.title, row, opening_col)
            ws.cell(row, opening_col).value = round(closing.value, 2)
            facts["carried"].append({"cell": opening, "label": _label(ws, row, opening_col, False),
                                     "from": closing.key, "amount": round(closing.value, 2)})
            for col in input_cols:
                if not _formula_text(ws.cell(row, col).value):
                    ws.cell(row, col).value = 0
                    reset.append(cell_key(ws.title, row, col))
        facts["ppe"].append({"sheet": ws.title, "reset": reset})


def _ppe_header(ws: Worksheet):
    """Row with "Opening" and "Closing" headings, plus "Addition"/"Deletion" columns."""
    for row in ws.iter_rows(max_row=min(ws.max_row, 30)):
        texts = {c.column: str(c.value).lower() for c in row if isinstance(c.value, str)}
        opening = next((col for col, t in texts.items() if "opening" in t), None)
        closing = next((col for col, t in texts.items() if "closing" in t), None)
        if opening and closing:
            inputs = [col for col, t in texts.items() if re.search(r"addition|deletion|deduction|sale", t)]
            return row[0].row, opening, closing, inputs
    return None


# ---------------------------------------------------------------- other facts for review

def _keep_manual(base: LearnResult, current_cols, facts):
    """Typed amounts in this year's columns stay as they were; the CA confirms or changes them."""
    carried = {c["cell"] for c in facts["carried"]}
    reset = {cell for p in facts["ppe"] for cell in p["reset"]}
    for key, cell in base.cells.items():
        if cell.category != MANUAL or key in carried or key in reset or not cell.value:
            continue
        sheet, address = _split(key)
        col = _row_col(address)[1]
        has_years = bool(current_cols.get(sheet))
        if has_years and col not in current_cols[sheet]:
            continue  # not a current-year column (e.g. a check column)
        if "rate" in cell.label.lower() or "%" in cell.label:
            continue  # rates carry over unchanged
        facts["manual"].append({"cell": key, "label": cell.label, "formula": cell.formula, "value": round(cell.value, 2)})


def link_suggestions(manual_cells: list[dict], tb: TrialBalance) -> list[dict]:
    """Typed amounts that equal TB ledgers (e.g. insurance =3540+6080). Each suggestion has a
    formula template that links them instead: ={TB:<Shop Insurance key>|Dr}+{TB:<Vehicle ...>|Dr}."""
    by_amount = {}
    for row in tb.balances:
        for side, amount in (("Dr", row.dr), ("Cr", row.cr)):
            if amount:
                by_amount.setdefault(amount, (row.key, side))
    suggestions = []
    for m in manual_cells:
        if not m["matches_ledgers"] or "opening" in m["label"].lower():
            continue  # an opening balance matching this year's TB is a coincidence, not a link
        pieces, ledgers = [], []
        for token in fx.tokens(m["formula"] or f"={m['value']}"):
            is_number = token.type == Token.OPERAND and token.subtype == Token.NUMBER
            match = by_amount.get(Decimal(token.value).quantize(Decimal("0.01"))) if is_number else None
            if match and match[0] not in ledgers:
                pieces.append(f"{{TB:{match[0]}|{match[1]}}}")
                ledgers.append(match[0])
            else:
                pieces.append(token.value)
        if ledgers:
            suggestions.append({"cell": m["cell"], "label": m["label"], "formula": m["formula"], "value": m["value"],
                                "template": "=" + "".join(pieces), "ledgers": ledgers})
    return suggestions


def _closing_stock(base: LearnResult) -> dict | None:
    """Last year's closing stock: the current-year cell that reads the Closing Stock line."""
    memo = {r.key for r in base.tb.rows if r.memo}
    for key, cell in base.cells.items():
        if cell.category == TB_LINKED and any(link.ledger in memo for link in cell.tb_links):
            return {"cell": key, "amount": round(cell.value, 2)}
    return None


def _replace_tb_sheet(target: Worksheet, source: Worksheet):
    """Put the new TB into the workbook's TB sheet (same sheet name, so other formulas still
    point at it), with Tally's formatting."""
    for merged in list(target.merged_cells.ranges):
        target.unmerge_cells(str(merged))
    target.delete_rows(1, target.max_row)
    for row in source.iter_rows():
        for c in row:
            if c.value is None and not c.has_style:
                continue
            new = target.cell(c.row, c.column, c.value)
            if c.has_style:
                new.font, new.border, new.fill = copy(c.font), copy(c.border), copy(c.fill)
                new.alignment, new.number_format = copy(c.alignment), c.number_format
    for merged in source.merged_cells.ranges:
        target.merge_cells(str(merged))
    for letter, dim in source.column_dimensions.items():
        target.column_dimensions[letter].width = dim.width


# ---------------------------------------------------------------- helpers

def _split(key: str) -> tuple[str, str]:
    sheet, address = key.rsplit("!", 1)
    return (sheet[1:-1].replace("''", "'") if sheet.startswith("'") else sheet), address


def _row_col(address: str) -> tuple[int, int]:
    letters, digits = re.match(r"([A-Z]+)(\d+)", address).groups()
    col = 0
    for ch in letters:
        col = col * 26 + ord(ch) - 64
    return int(digits), col


def _fy(start: int) -> str:
    return f"{start}-{(start + 1) % 100:02d}"


def _tb_summary(tb: TrialBalance) -> str:
    if tb.total_dr == tb.total_cr:
        return f'Sheet "{tb.sheet_name}". Debit and credit both ₹{format_inr(tb.total_dr)}.'
    return f'Sheet "{tb.sheet_name}". Debit ₹{format_inr(tb.total_dr)} and credit ₹{format_inr(tb.total_cr)} do not match.'
