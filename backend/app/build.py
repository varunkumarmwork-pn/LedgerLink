"""BUILD MODE: financial statements from a Tally TB alone, in a default format.

Every ledger with a balance is put on a statement line by its Tally group (templates.py).
Only the sheets the user ticked are written; every figure is a live formula to the TB:
  Note row         =TB!C19-TB!B19            (one ledger; Cr - Dr for liabilities and income)
  Statement line   ='BS Schedules'!C12        (the note total), or the ledgers directly when
                                               the notes sheet was not ticked
  Profit           income - expenses; it goes to capital (non-corporate) or reserves (company)
Ledgers no line takes are left out, so review shows them red as not used.
"""
from copy import copy
from datetime import datetime
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Border, Font, Side
from openpyxl.utils import column_index_from_string

from app.learn import Progress, _fy_start_from_period
from app.tally_tb import TBRow, TrialBalance, format_inr, parse_trial_balance
from app.templates import CAPITAL_LABELS, LINES, is_company, line_for, template_for

SHEETS = ["Balance Sheet", "Profit & Loss", "BS Schedules", "P&L Notes", "PPE", "Accounting policies"]
INDIAN = r'[>=10000000]##\,##\,##\,##0.00;[>=100000]##\,##\,##0.00;##,##0.00'  # 1,36,74,167.85 in Excel
FONT, BOLD, TITLE = Font(name="Arial", size=10), Font(name="Arial", size=10, bold=True), Font(name="Arial", size=13, bold=True)
THIN = Side(style="thin")


def build_workbook(tb_path: str | Path, entity_type: str, client_name: str, sheets: list[str],
                   out_path: str | Path, progress: Progress | None = None) -> dict:
    """Write the ticked sheets plus the TB to out_path. Returns what was built (JSON-ready)."""
    progress = progress or Progress()
    sheets = [s for s in SHEETS if s in sheets]
    if not sheets:
        raise ValueError("Tick at least one sheet to build.")
    tb = parse_trial_balance(tb_path)
    progress.step_done(0, _tb_summary(tb))

    lines, unmapped = assign_lines(tb, entity_type)
    mapped = sum(len(rows) for rows in lines.values())
    progress.step_done(1, f"{mapped} ledgers put on {len(lines)} statement lines; {len(unmapped)} with no default line.")

    writer = _Writer(tb, entity_type, client_name, sheets, lines, openpyxl.load_workbook(tb_path)[tb.sheet_name])
    writer.write()
    writer.wb.save(out_path)
    progress.step_done(2, f"Wrote {', '.join(sheets)} and the TB, with live formulas to the TB.")
    return {"template": template_for(entity_type)["name"], "sheets": sheets,
            "lines": {line: [r.key for r in rows] for line, rows in lines.items()},
            "unmapped": [r.key for r in unmapped]}


def assign_lines(tb: TrialBalance, entity_type: str) -> tuple[dict[str, list[TBRow]], list[TBRow]]:
    """{line id: [TB rows]} for every balance, and the rows no line takes."""
    lines: dict[str, list[TBRow]] = {}
    unmapped = []
    for row in tb.balances:
        line = line_for([row.name] + list(reversed(row.path)), entity_type)
        if line is None:
            unmapped.append(row)
            continue
        lines.setdefault(line, []).append(row)
        if row.memo and line == "inventories":
            # Tally's closing stock is an asset and also reduces this year's cost of goods
            lines.setdefault("inventory_change", []).append(row)
    return lines, unmapped


class _Writer:
    def __init__(self, tb, entity_type, client_name, sheets, lines, tb_source):
        self.tb, self.entity_type, self.client, self.sheets, self.lines = tb, entity_type, client_name, sheets, lines
        self.template = template_for(entity_type)
        start = _fy_start_from_period(tb.period) or datetime.now().year
        self.fy, self.prior_fy = _fy(start), _fy(start - 1)
        self.year_end = f"31 March {start + 1}"
        self.tb_ref = tb.sheet_name if tb.sheet_name.replace("_", "").isalnum() else f"'{tb.sheet_name}'"
        self.notes: dict[str, str] = {}  # line id -> cell with its total (a note, or PPE)
        self.numbers: dict[str, int] = {}  # line id -> note number
        self.profit: str | None = None  # formula text for this year's profit

        self.wb = openpyxl.Workbook()
        self.wb.remove(self.wb.active)
        order = [s for s in SHEETS if s in sheets and s != "Accounting policies"] + [tb.sheet_name] + \
                (["Accounting policies"] if "Accounting policies" in sheets else [])
        for title in order:
            self.wb.create_sheet(title)
        _copy_sheet(tb_source, self.wb[tb.sheet_name])

    # ------------------------------------------------------------ what goes on each line

    def items(self, line: str) -> list[tuple[str, str]]:
        """(label, formula text) for each TB row on a line."""
        side = LINES[line][1]
        result = []
        for row in self.lines.get(line, []):
            dr, cr = f"{self.tb_ref}!{self.tb.debit_column}{row.row}", f"{self.tb_ref}!{self.tb.credit_column}{row.row}"
            if row.memo:  # Tally's closing stock line: an asset, and a reduction of expenses
                result.append(("Less: " + row.name if side == "expense" else row.name, f"-{cr}" if side == "expense" else cr))
            elif side in ("liability", "income"):
                result.append((row.name, f"{cr}-{dr}"))
            else:
                result.append((row.name, f"{dr}-{cr}"))
        return result

    def owner_line(self) -> str:
        return "reserves" if is_company(self.entity_type) else "capital"

    def label(self, line: str) -> str:
        if line == "capital":
            return CAPITAL_LABELS.get(self.entity_type, LINES[line][0])
        return LINES[line][0]

    def has_content(self, line: str) -> bool:
        return bool(self.lines.get(line)) or line == self.owner_line() or \
            (line in ("ppe", "depreciation") and "PPE" in self.sheets and self.lines.get("ppe"))

    def extra(self, line: str) -> list[tuple[str, str]]:
        """Rows that are not TB ledgers: profit into capital/reserves, depreciation from PPE."""
        if line == self.owner_line():
            return [("Add: Profit for the year", self.profit_formula())]
        if line == "depreciation" and "PPE" in self.sheets and self.lines.get("ppe"):
            return [("Depreciation as per PPE schedule", self.ppe_depreciation)]
        return []

    def value_formula(self, line: str) -> str:
        """A statement line: its note total, or the TB ledgers directly."""
        if line in self.notes:
            return f"={self.notes[line]}"
        parts = [f"({f})" for _, f in self.items(line) + self.extra(line)]
        return "=" + ("+".join(parts) if parts else "0")

    def profit_formula(self) -> str:
        if self.profit:
            return self.profit
        pl = [line for kind, *rest in self.template["profit_and_loss"] if kind == "line" for line in rest] + ["tax"]
        income = [f"({f})" for line in pl if LINES[line][1] == "income" for _, f in self.items(line)]
        expense = [f"({f})" for line in pl if LINES[line][1] == "expense" for _, f in self.items(line) + self.extra(line)]
        return f"({'+'.join(income) or '0'})-({'+'.join(expense) or '0'})"

    # ------------------------------------------------------------ writing

    def write(self):
        self.number_notes()
        if "PPE" in self.sheets and self.lines.get("ppe"):
            self.write_ppe(self.wb["PPE"])
        if "P&L Notes" in self.sheets:
            self.write_notes(self.wb["P&L Notes"], "Notes forming part of the Statement of Profit and Loss",
                             self.pl_lines())
        if "Profit & Loss" in self.sheets:
            self.write_profit_and_loss(self.wb["Profit & Loss"])
        if "BS Schedules" in self.sheets:
            self.write_notes(self.wb["BS Schedules"], "Notes forming part of the Balance Sheet", self.bs_lines())
        if "Balance Sheet" in self.sheets:
            self.write_balance_sheet(self.wb["Balance Sheet"])
        if "Accounting policies" in self.sheets:
            self.write_policies(self.wb["Accounting policies"])
        if "PPE" in self.sheets and not self.lines.get("ppe"):
            self.wb["PPE"]["B2"] = "No fixed asset ledgers in the trial balance."

    def bs_lines(self) -> list[str]:
        return [rest[0] for kind, *rest in self.template["balance_sheet"] if kind == "line" and self.has_content(rest[0])]

    def pl_lines(self) -> list[str]:
        lines = [rest[0] for kind, *rest in self.template["profit_and_loss"] if kind == "line"] + ["tax"]
        return [line for line in lines if self.has_content(line)]

    def number_notes(self):
        for n, line in enumerate(self.bs_lines() + self.pl_lines(), start=1):
            self.numbers[line] = n

    def write_notes(self, ws, title, lines):
        _widths(ws, {"A": 3, "B": 52, "C": 18, "D": 18})
        _put(ws, 1, 2, self.client, TITLE)
        _put(ws, 2, 2, title, BOLD)
        row = 4
        for line in lines:
            if line == "ppe" and "ppe" in self.notes:
                continue  # the PPE sheet is this note
            _put(ws, row, 2, f"Note {self.numbers[line]} – {self.label(line)}", BOLD)
            row += 1
            for col, text in ((2, "Particulars"), (3, f"FY {self.fy} (₹)"), (4, f"FY {self.prior_fy} (₹)")):
                _put(ws, row, col, text, BOLD, border=True, align="center" if col > 2 else None)
            first = row + 1
            for label, formula in self.items(line) + self.extra(line):
                row += 1
                _put(ws, row, 2, label, FONT, border=True)
                _put(ws, row, 3, f"={formula}", FONT, border=True, number=True)
                _put(ws, row, 4, None, FONT, border=True, number=True)
            row += 1
            _put(ws, row, 2, f"Total {self.label(line)}", BOLD, border=True)
            _put(ws, row, 3, f"=SUM(C{first}:C{row - 1})", BOLD, border=True, number=True)
            _put(ws, row, 4, None, BOLD, border=True, number=True)
            sheet = f"'{ws.title}'" if not ws.title.replace("_", "").isalnum() else ws.title
            self.notes[line] = f"{sheet}!C{row}"
            row += 2

    def write_statement(self, ws, title, rows) -> dict[str, int]:
        """Headings, lines and totals of one statement. Returns {line id or total text: row}."""
        _widths(ws, {"A": 3, "B": 52, "C": 8, "D": 18, "E": 18})
        _put(ws, 1, 2, self.client, TITLE)
        _put(ws, 2, 2, title, BOLD)
        for col, text in ((2, "Particulars"), (3, "Note"), (4, f"FY {self.fy} (₹)"), (5, f"FY {self.prior_fy} (₹)")):
            _put(ws, 4, col, text, BOLD, border=True, align="center" if col > 2 else None)
        where, row = {}, 5
        for kind, *rest in rows:
            if kind == "line" and not self.has_content(rest[0]):
                continue
            if kind == "heading":
                _put(ws, row, 2, rest[0], BOLD)
            elif kind == "line":
                line = rest[0]
                _put(ws, row, 2, self.label(line), FONT, border=True)
                _put(ws, row, 3, self.numbers.get(line), FONT, border=True, align="center")
                _put(ws, row, 4, self.value_formula(line), FONT, border=True, number=True)
                _put(ws, row, 5, None, FONT, border=True, number=True)
                where[line] = row
            else:  # total of the lines above that are shown
                text, members = rest
                refs = [f"D{where[m]}" for m in members if m in where]
                _put(ws, row, 2, text, BOLD, border=True)
                _put(ws, row, 4, "=" + ("+".join(refs) if refs else "0"), BOLD, border=True, number=True)
                _put(ws, row, 5, None, BOLD, border=True, number=True)
                where[text] = row
            row += 1
        return where

    def write_profit_and_loss(self, ws):
        where = self.write_statement(ws, f"Statement of Profit and Loss for the year ended {self.year_end}",
                                     self.template["profit_and_loss"])
        row = max(where.values()) + 1
        _put(ws, row, 2, "Profit before Tax", BOLD, border=True)
        _put(ws, row, 4, f"=D{where['Total Income']}-D{where['Total Expenses']}", BOLD, border=True, number=True)
        before_tax = row
        if self.has_content("tax"):
            row += 1
            _put(ws, row, 2, self.label("tax"), FONT, border=True)
            _put(ws, row, 3, self.numbers.get("tax"), FONT, border=True, align="center")
            _put(ws, row, 4, self.value_formula("tax"), FONT, border=True, number=True)
        row += 1
        _put(ws, row, 2, "Profit for the Year", BOLD, border=True)
        _put(ws, row, 4, f"=D{before_tax}-D{row - 1}" if row - 1 != before_tax else f"=D{before_tax}",
             BOLD, border=True, number=True)
        self.profit = f"'Profit & Loss'!D{row}"

    def write_balance_sheet(self, ws):
        self.write_statement(ws, f"Balance Sheet as at {self.year_end}", self.template["balance_sheet"])

    def write_ppe(self, ws):
        """Fixed asset ledgers with a depreciation rate for the CA to fill in (0 until then)."""
        _widths(ws, {"A": 3, "B": 36, "C": 18, "D": 10, "E": 16, "F": 18})
        _put(ws, 1, 2, self.client, TITLE)
        _put(ws, 2, 2, f"Note {self.numbers.get('ppe', '')} – Property, Plant and Equipment", BOLD)
        for col, text in enumerate(["Particulars", "Balance as per TB", "Rate", "Depreciation", "Closing WDV"], start=2):
            _put(ws, 3, col, text, BOLD, border=True, align="center" if col > 2 else None)
        row = 3
        for label, formula in self.items("ppe"):
            row += 1
            _put(ws, row, 2, label, FONT, border=True)
            _put(ws, row, 3, f"={formula}", FONT, border=True, number=True)
            cell = _put(ws, row, 4, 0, FONT, border=True)
            cell.number_format = "0%"
            _put(ws, row, 5, f"=ROUND(C{row}*D{row},2)", FONT, border=True, number=True)
            _put(ws, row, 6, f"=C{row}-E{row}", FONT, border=True, number=True)
        total = row + 1
        _put(ws, total, 2, "Total", BOLD, border=True)
        for col in "CEF":
            _put(ws, total, column_index_from_string(col), f"=SUM({col}4:{col}{row})", BOLD, border=True, number=True)
        self.notes["ppe"] = f"PPE!F{total}"
        self.ppe_depreciation = f"PPE!E{total}"

    def write_policies(self, ws):
        _widths(ws, {"A": 4, "B": 100})
        _put(ws, 1, 1, self.client, TITLE)
        _put(ws, 2, 1, "Significant Accounting Policies", BOLD)
        row = 4
        for letter, (heading, text) in zip("abcdefghij", self.template["policies"]):
            _put(ws, row, 1, f"{letter}.", BOLD)
            _put(ws, row, 2, heading.upper(), BOLD)
            cell = _put(ws, row + 1, 2, text, FONT)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            row += 3


# ---------------------------------------------------------------- helpers

def _put(ws, row, col, value, font, border=False, number=False, align=None):
    cell = ws.cell(row, col, value)
    cell.font = font
    if border:
        cell.border = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
    if number:
        cell.number_format = INDIAN
    if align:
        cell.alignment = Alignment(horizontal=align)
    return cell


def _widths(ws, widths: dict[str, float]):
    for letter, width in widths.items():
        ws.column_dimensions[letter].width = width
    ws.sheet_view.showGridLines = False


def _copy_sheet(source, target):
    for row in source.iter_rows():
        for c in row:
            if c.value is None and not c.has_style:
                continue
            new = target.cell(c.row, c.column, c.value)
            if c.has_style:
                new.font, new.border, new.fill = copy(c.font), copy(c.border), copy(c.fill)
                new.alignment, new.number_format = copy(c.alignment), c.number_format
    for merged in source.merged_cells.ranges:
        target.merge_cells(str(merged))
    for letter, dim in source.column_dimensions.items():
        target.column_dimensions[letter].width = dim.width


def _fy(start: int) -> str:
    return f"{start}-{(start + 1) % 100:02d}"


def _tb_summary(tb: TrialBalance) -> str:
    if tb.total_dr == tb.total_cr:
        return f'Sheet "{tb.sheet_name}". Debit and credit both ₹{format_inr(tb.total_dr)}.'
    return f'Sheet "{tb.sheet_name}". Debit ₹{format_inr(tb.total_dr)} and credit ₹{format_inr(tb.total_cr)} do not match.'
