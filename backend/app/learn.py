"""LEARN MODE: understand a client's finished workbook so next year's can be rebuilt from it.

For every numeric cell outside the TB sheet we work out where its figure comes from:
  tb-linked   formula reads the TB, e.g. =TB!B48 (Bank of Baroda, Dr)
  internal    formula reads other cells only, e.g. ='BS Schedules'!C58 or =SUM(C55:C57)
  manual      typed number, or a formula of constants only, e.g. =3540+6080
  prior-year  anything in a column headed with an earlier FY (mostly typed last-year figures)

TB references are stored by ledger NAME (group path + name), never by row, because the
rows move every year when Tally adds or drops ledgers.
"""
import re
from collections import Counter, defaultdict
from dataclasses import dataclass, field
from datetime import datetime
from decimal import Decimal
from pathlib import Path

import openpyxl
from openpyxl.formula.tokenizer import Token
from openpyxl.utils import column_index_from_string
from openpyxl.worksheet.worksheet import Worksheet

from app import formula as fx
from app.tally_tb import TrialBalance, format_inr, parse_trial_balance

TB_LINKED, INTERNAL, MANUAL, PRIOR_YEAR = "tb-linked", "internal", "manual", "prior-year"
FY_IN_TEXT = re.compile(r"(?<!\d)(20\d{2})\s*[-–]\s*(\d{2})(?!\d)")  # "FY 2025-26 (₹)"
TALLY_DATE = re.compile(r"\d{1,2}-[A-Za-z]{3}-\d{2,4}")  # "31-Mar-26"
NOTE_HEADINGS = {"note", "notes", "note no", "note no."}  # columns of note numbers, not amounts
HALF_PAISA = 0.005


@dataclass
class TBLink:
    ledger: str  # TBRow.key, e.g. "Current Assets > Bank Accounts > Bank of Baroda"
    side: str  # "Dr" or "Cr"


@dataclass
class CellInfo:
    key: str  # e.g. "'BS Schedules'!C55"
    label: str  # plain words for the row, e.g. "Bank of Baroda"
    formula: str | None
    value: float  # recalculated by us (Excel's saved value if we could not)
    category: str
    financial_year: str | None  # from the column heading; None if the column has none
    tb_links: list[TBLink] = field(default_factory=list)
    template: str | None = None  # formula with each TB reference written as {TB:<ledger key>|<side>}


@dataclass
class LearnResult:
    financial_year: str | None
    tb: TrialBalance
    cells: dict[str, CellInfo]
    precedents: dict[str, list[str]]  # dependency graph: formula cell -> every cell it reads
    unused_ledgers: list[dict]
    manual_cells: list[dict]
    inconsistencies: list[dict]
    warnings: list[str]
    year_columns: dict[str, dict[int, int]] = field(default_factory=dict)  # sheet -> {column: FY start year}

    def mapping(self) -> dict:
        """What we store per client: TB ledgers by name and how each cell is built from them."""
        ledgers = {r.key: {"name": r.name, "group_path": r.group_path, "kind": r.kind,
                           "dr": float(r.dr), "cr": float(r.cr), "used_in": []} for r in self.tb.rows}
        for cell in self.cells.values():
            for link in cell.tb_links:
                ledgers[link.ledger]["used_in"].append({"cell": cell.key, "side": link.side})
        cells = {
            c.key: {"label": c.label, "category": c.category, "financial_year": c.financial_year,
                    "formula": c.formula, "template": c.template, "value": round(c.value, 2),
                    "links": [{"ledger": l.ledger, "side": l.side} for l in c.tb_links]}
            for c in self.cells.values()
        }
        return {"financial_year": self.financial_year, "tb_sheet": self.tb.sheet_name,
                "ledgers": ledgers, "cells": cells, "precedents": self.precedents,
                "year_columns": {sheet: {str(col): year for col, year in cols.items()}
                                 for sheet, cols in self.year_columns.items() if cols}}

    def report(self) -> dict:
        return {
            "financial_year": self.financial_year,
            "tb": {"sheet": self.tb.sheet_name, "company": self.tb.company,
                   "total_dr": float(self.tb.total_dr), "total_cr": float(self.tb.total_cr),
                   "balanced": self.tb.is_balanced, "problems": self.tb.problems},
            "cell_counts": dict(Counter(c.category for c in self.cells.values())),
            "unused_ledgers": self.unused_ledgers,
            "manual_cells": self.manual_cells,
            "inconsistencies": self.inconsistencies,
            "warnings": self.warnings,
        }


class Progress:
    """Receives progress while a workbook is learnt. Does nothing by default; the
    background job overrides it to show the steps on the Processing screen."""

    # The four steps, in order: 0 read TB, 1 find sheets, 2 trace links, 3 check totals.
    def step_done(self, index: int, detail: str):
        pass

    def traced(self, cell: str, formula: str, ledgers: str):
        pass


def learn_workbook(path: str | Path, progress: Progress | None = None) -> LearnResult:
    return _Learner(path, progress or Progress()).run()


class _Learner:
    def __init__(self, path: str | Path, progress: Progress):
        self.progress = progress
        # Loaded twice: once for formulas, once for the values Excel saved with them.
        self.wb = openpyxl.load_workbook(path)
        self.saved = openpyxl.load_workbook(path, data_only=True)
        self.tb = parse_trial_balance(self.saved)
        progress.step_done(0, _tb_summary(self.tb))
        self.tb_rows = {r.row: r for r in self.tb.rows}
        self.tb_by_key = {r.key: r for r in self.tb.rows}
        self.tb_sides = {column_index_from_string(self.tb.debit_column): "Dr",
                         column_index_from_string(self.tb.credit_column): "Cr"}
        self.sheets = [ws for ws in self.wb.worksheets if ws.title != self.tb.sheet_name]
        self.column_years = {ws.title: _column_years(ws) for ws in self.sheets}
        all_years = [y for years in self.column_years.values() for y in years.values()]
        self.current_year = _fy_start_from_period(self.tb.period) or max(all_years, default=None)
        progress.step_done(1, ", ".join(ws.title for ws in self.wb.worksheets))

        self.values: dict[str, object] = {}
        self.calculating: set[str] = set()
        self.inconsistencies: list[dict] = []
        self.warnings: list[str] = []

    def run(self) -> LearnResult:
        cells, precedents = {}, {}
        for ws in self.sheets:
            note_columns = _note_columns(ws)
            filled = [c for row in ws.iter_rows() for c in row if c.value is not None]
            for c in filled:
                formula = _formula_text(c.value)
                if formula:
                    key = fx.cell_key(ws.title, c.row, c.column)
                    precedents[key] = self._precedents(key, ws.title, formula)
                if c.column not in note_columns:
                    info = self._classify(ws, c.row, c.column, formula)
                    if info:
                        cells[info.key] = info
        tb_linked = sum(1 for c in cells.values() if c.category == TB_LINKED)
        self.progress.step_done(2, f"{len(precedents)} formulas followed; {tb_linked} read the trial balance.")

        self._check_year_columns(cells)
        result = LearnResult(
            financial_year=_fy_label(self.current_year),
            tb=self.tb,
            cells=cells,
            precedents=precedents,
            unused_ledgers=self._unused_ledgers(cells),
            manual_cells=self._manual_cells(cells),
            inconsistencies=self.inconsistencies,
            warnings=self.warnings,
            year_columns=self.column_years,
        )
        self.progress.step_done(3, f"{_count(len(result.unused_ledgers), 'TB ledger')} not used, "
                                   f"{_count(len(result.manual_cells), 'typed amount')}, "
                                   f"{_count(len(result.inconsistencies), 'inconsistency', 'inconsistencies')}.")
        return result

    # ------------------------------------------------------------ classifying cells

    def _classify(self, ws: Worksheet, row: int, col: int, formula: str | None) -> CellInfo | None:
        key = fx.cell_key(ws.title, row, col)
        value = self.value_of(ws.title, row, col)
        if not _is_number(value):
            return None  # labels, headings, dates

        year = self.column_years[ws.title].get(col)
        is_prior_year = year is not None and self.current_year is not None and year < self.current_year
        label = _label(ws, row, col, has_year_heading=year is not None)
        tb_links, template = self._tb_links(key, label, ws.title, formula) if formula else ([], None)

        if is_prior_year:
            category = PRIOR_YEAR
        elif tb_links:
            category = TB_LINKED
        elif formula and fx.references(formula, ws.title):
            category = INTERNAL
        else:
            category = MANUAL

        if formula:
            self._compare_with_saved(key, label, ws.title, row, col, value)
        if category == TB_LINKED:
            names = dict.fromkeys(self.tb_by_key[link.ledger].name for link in tb_links)
            self.progress.traced(key, formula, ", ".join(names))
        return CellInfo(key=key, label=label, formula=formula, value=float(value), category=category,
                        financial_year=_fy_label(year), tb_links=tb_links, template=template)

    def _tb_links(self, key: str, label: str, sheet: str, formula: str) -> tuple[list[TBLink], str | None]:
        """Resolve every TB reference in a formula to ledger names. Returns (links, template)."""
        links, pieces = [], []
        for token in fx.tokens(formula):
            ref = fx.parse_ref(token.value, sheet) if token.subtype == Token.RANGE else None
            if ref is None or ref.sheet != self.tb.sheet_name:
                pieces.append(token.value)
                continue
            if ref.min_row is None or ref.min_col is None:
                self.warnings.append(f"{key}: {token.value} refers to whole TB rows or columns; "
                                     f"it cannot be mapped to ledgers.")
                pieces.append(token.value)
                continue
            covered = [self._tb_link(key, label, r, c, report=ref.is_single_cell)
                       for r, c in ref.cells(ref.max_row, ref.max_col)]
            links += [link for link in covered if link]
            ends = [covered[0], covered[-1]] if not ref.is_single_cell else [covered[0]]
            pieces.append(":".join(f"{{TB:{e.ledger}|{e.side}}}" for e in ends) if all(ends) else token.value)
        return links, ("=" + "".join(pieces) if links else None)

    def _tb_link(self, key: str, label: str, row: int, col: int, report: bool) -> TBLink | None:
        tb_row, side = self.tb_rows.get(row), self.tb_sides.get(col)
        if tb_row and side:
            return TBLink(tb_row.key, side)
        if report:
            where = fx.cell_key(self.tb.sheet_name, row, col)
            self._inconsistent(key, label, f"links {where}, which is not a ledger's Dr or Cr amount in the TB.")
        return None

    # ------------------------------------------------------------ dependency graph and recalculation

    def _precedents(self, key: str, sheet: str, formula: str) -> list[str]:
        found = []
        for text, ref in fx.references(formula, sheet):
            if ref is None or ref.sheet not in self.wb.sheetnames:
                self.warnings.append(f"{key}: cannot follow reference {text}.")
                continue
            ws = self.wb[ref.sheet]
            found += [fx.cell_key(ref.sheet, r, c) for r, c in ref.cells(ws.max_row, ws.max_column)]
        return found

    def value_of(self, sheet: str, row: int, col: int):
        """A cell's value, recalculating formulas from their precedents."""
        key = fx.cell_key(sheet, row, col)
        if key in self.values:
            return self.values[key]
        raw = self.wb[sheet].cell(row, col).value
        formula = _formula_text(raw)
        if not formula:
            self.values[key] = raw
            return raw
        if key in self.calculating:
            raise fx.FormulaError(f"circular reference at {key}")
        self.calculating.add(key)
        try:
            value = fx.evaluate(formula, lambda text: self._lookup(text, sheet))
        except fx.FormulaError as error:
            value = self.saved[sheet].cell(row, col).value
            self.warnings.append(f"{key}: could not recalculate {formula} ({error}); used Excel's saved value.")
        finally:
            self.calculating.discard(key)
        self.values[key] = value
        return value

    def _lookup(self, text: str, sheet: str):
        ref = fx.parse_ref(text, sheet)
        if ref is None or ref.sheet not in self.wb.sheetnames:
            raise fx.FormulaError(f"cannot follow reference {text}")
        if ref.is_single_cell:
            return self.value_of(ref.sheet, ref.min_row, ref.min_col)
        ws = self.wb[ref.sheet]
        return [self.value_of(ref.sheet, r, c) for r, c in ref.cells(ws.max_row, ws.max_column)]

    def _compare_with_saved(self, key: str, label: str, sheet: str, row: int, col: int, value):
        saved = self.saved[sheet].cell(row, col).value
        if _is_number(saved) and abs(saved - value) > HALF_PAISA:
            self._inconsistent(key, label, f"Excel's saved value {format_inr(saved)} differs from the recalculated "
                                           f"{format_inr(value)}. Open and save the file in Excel, then upload again.")

    # ------------------------------------------------------------ report

    def _check_year_columns(self, cells: dict[str, CellInfo]):
        """This year's and last year's formula on the same row should read the same lines.

        Example (Venus, Balance Sheet row 21): FY 2025-26 reads 'BS Schedules'!C74 (one line),
        FY 2024-25 reads 'BS Schedules'!D76 (the total). Columns may differ (PPE closing vs
        opening), so only the rows referenced and the formula shape are compared.
        """
        if self.current_year is None:
            return
        for ws in self.sheets:
            years = self.column_years[ws.title]
            current_cols = sorted(c for c, y in years.items() if y == self.current_year)
            prior_cols = sorted(c for c, y in years.items() if y == self.current_year - 1)
            for current_col, prior_col in zip(current_cols, prior_cols):
                for row in range(1, ws.max_row + 1):
                    self._compare_year_pair(ws, row, current_col, prior_col)

    def _compare_year_pair(self, ws: Worksheet, row: int, current_col: int, prior_col: int):
        current = _formula_text(ws.cell(row, current_col).value)
        prior = _formula_text(ws.cell(row, prior_col).value)
        if not (current and prior and fx.references(current, ws.title) and fx.references(prior, ws.title)):
            return  # last year typed in: nothing to compare
        if self._shape(current, ws.title, self.current_year) == self._shape(prior, ws.title, self.current_year - 1):
            return
        key = fx.cell_key(ws.title, row, current_col)
        label = _label(ws, row, current_col, has_year_heading=True)
        self._inconsistent(key, label,
                           f"FY {_fy_label(self.current_year)} uses {current} but FY "
                           f"{_fy_label(self.current_year - 1)} uses {prior}. They read different lines; "
                           f"one of them is probably wrong.")

    def _shape(self, formula: str, sheet: str, own_year: int) -> list[str]:
        """Formula with each reference reduced to sheet + rows + which year its column holds."""
        shape = []
        for token in fx.tokens(formula):
            ref = fx.parse_ref(token.value, sheet) if token.subtype == Token.RANGE else None
            if ref is None:
                shape.append(token.value.upper())
                continue
            col_year = self.column_years.get(ref.sheet, {}).get(ref.min_col)
            year_tag = "" if col_year is None else f"year{col_year - own_year:+d}"
            shape.append(f"{ref.sheet}!R{ref.min_row}:R{ref.max_row}{year_tag}")
        return shape

    def _unused_ledgers(self, cells: dict[str, CellInfo]) -> list[dict]:
        """TB balances no formula reads. A reference to a group covers every ledger under it."""
        covered = set()
        for cell in cells.values():
            for link in cell.tb_links:
                covered |= {r.key for r in self.tb_by_key[link.ledger].with_descendants()}
        typed = self._typed_amounts(cells)
        return [
            {"ledger": b.key, "name": b.name, "group_path": b.group_path, "dr": float(b.dr), "cr": float(b.cr),
             # the amount may have been typed in somewhere instead of linked
             "typed_in": sorted({cell for amount in (b.dr, b.cr) if amount for cell in typed.get(amount, [])})}
            for b in self.tb.balances if b.key not in covered
        ]

    def _manual_cells(self, cells: dict[str, CellInfo]) -> list[dict]:
        by_amount = defaultdict(list)
        for b in self.tb.balances:
            for amount in (b.dr, b.cr):
                if amount:
                    by_amount[amount].append(b.name)
        return [
            {"cell": c.key, "label": c.label, "formula": c.formula, "value": round(c.value, 2),
             # typed figures that equal a TB ledger: probably should be linked instead
             "matches_ledgers": [name for amount in _typed_numbers(c) for name in by_amount.get(amount, [])]}
            for c in cells.values() if c.category == MANUAL
        ]

    def _typed_amounts(self, cells: dict[str, CellInfo]) -> dict[Decimal, list[str]]:
        typed = defaultdict(list)
        for c in cells.values():
            if c.category == MANUAL:
                for amount in _typed_numbers(c):
                    typed[amount].append(c.key)
        return typed

    def _inconsistent(self, key: str, label: str, message: str):
        self.inconsistencies.append({"cell": key, "label": label, "message": message})


# ---------------------------------------------------------------- helpers

def _tb_summary(tb: TrialBalance) -> str:
    if tb.total_dr == tb.total_cr:
        return f'Sheet "{tb.sheet_name}". Debit and credit both ₹{format_inr(tb.total_dr)}.'
    return (f'Sheet "{tb.sheet_name}". Debit ₹{format_inr(tb.total_dr)} and credit ₹{format_inr(tb.total_cr)} '
            f'do not match.')


def _count(n: int, singular: str, plural: str | None = None) -> str:
    return f"{n} {singular if n == 1 else plural or singular + 's'}"


def _formula_text(value) -> str | None:
    if isinstance(value, str) and value.startswith("=") and len(value) > 1:
        return value
    text = getattr(value, "text", None)  # openpyxl ArrayFormula
    return text if isinstance(text, str) else None


def _is_number(value) -> bool:
    return isinstance(value, (int, float)) and not isinstance(value, bool)


def _paise(value: float) -> Decimal:
    return Decimal(str(value)).quantize(Decimal("0.01"))


def _typed_numbers(cell: CellInfo) -> list[Decimal]:
    numbers = fx.constants(cell.formula) if cell.formula else [cell.value]
    return [_paise(n) for n in numbers if n]


def _fy_label(start_year: int | None) -> str | None:
    return None if start_year is None else f"{start_year}-{(start_year + 1) % 100:02d}"


def _fy_start_from_period(period: str | None) -> int | None:
    """"1-Apr-25 to 31-Mar-26" -> 2025 (the FY in which the period ends)."""
    dates = TALLY_DATE.findall(period or "")
    if not dates:
        return None
    for pattern in ("%d-%b-%y", "%d-%b-%Y"):
        try:
            end = datetime.strptime(dates[-1], pattern)
            return end.year - 1 if end.month <= 3 else end.year
        except ValueError:
            continue
    return None


def _column_years(ws: Worksheet) -> dict[int, int]:
    """Which FY each column holds, from headings like "FY 2024-25 (₹)" or a 31 March date."""
    votes = defaultdict(Counter)
    for row in ws.iter_rows():
        for c in row:
            if isinstance(c.value, str):
                match = FY_IN_TEXT.search(c.value)
                if match and int(match[2]) == (int(match[1]) + 1) % 100:
                    votes[c.column][int(match[1])] += 1
            elif isinstance(c.value, datetime) and (c.value.month, c.value.day) == (3, 31):
                votes[c.column][c.value.year - 1] += 1
    return {col: counter.most_common(1)[0][0] for col, counter in votes.items()}


def _note_columns(ws: Worksheet) -> set[int]:
    return {c.column for row in ws.iter_rows() for c in row
            if isinstance(c.value, str) and c.value.strip().lower() in NOTE_HEADINGS}


def _label(ws: Worksheet, row: int, col: int, has_year_heading: bool) -> str:
    """Plain-words name for a cell: the row's text, plus the column heading when the
    column is not simply a year (e.g. PPE "Car – Opening Balance")."""
    row_text = next((str(ws.cell(row, c).value).strip() for c in range(1, col)
                     if isinstance(ws.cell(row, c).value, str) and not _formula_text(ws.cell(row, c).value)), "")
    if has_year_heading:
        return row_text
    heading = next((str(ws.cell(r, col).value).strip() for r in range(row - 1, 0, -1)
                    if isinstance(ws.cell(r, col).value, str) and not _formula_text(ws.cell(r, col).value)), "")
    return f"{row_text} – {heading}" if row_text and heading else row_text or heading
