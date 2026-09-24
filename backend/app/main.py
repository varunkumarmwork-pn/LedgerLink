"""LedgerLink API. Start with:  uvicorn app.main:app --reload --port 8001  (from the backend folder).
The frontend (Vite) forwards /api to port 8001; 8000 is used by another app on this machine."""
import asyncio
import json
import logging
import os
import re
import threading
import zipfile
from contextlib import asynccontextmanager
from datetime import date, datetime
from pathlib import Path
from typing import Literal

from fastapi import BackgroundTasks, Depends, FastAPI, Form, HTTPException, UploadFile
from fastapi.responses import Response, StreamingResponse
from openpyxl.utils.exceptions import InvalidFileException
from pydantic import BaseModel, ConfigDict, Field
from sqlalchemy import select
from sqlalchemy.orm import Session, sessionmaker

from app import ai, jobs
from app.build import SHEETS as BUILD_SHEETS
from app.build import build_workbook
from app.db import create_schema, get_session, get_session_factory
from app.excel import learn_effective, workbook_bytes
from app.learn import learn_workbook
from app.models import Client, ClientWorkbook, EntityType
from app.review import apply_action, build_review
from app.rollforward import link_suggestions, roll_forward
from app.tally_tb import parse_trial_balance
from app.workbook_view import apply_edits, edited_cells, workbook_view

# Uploaded workbooks are kept here, one folder per client.
STORAGE_DIR = Path(os.environ.get("LEDGERLINK_STORAGE", Path(__file__).resolve().parent.parent / "storage"))
# Review status shown on the Clients screen while and after a workbook is read.
READING, NOT_READ, ROLLING, BUILDING = "Reading workbook", "Workbook not read", "Rolling forward", "Building"
logger = logging.getLogger("ledgerlink")


@asynccontextmanager
async def lifespan(app: FastAPI):
    create_schema()
    # Load the local language model now, so the first job does not wait for it.
    threading.Thread(target=ai.meaning.available, daemon=True).start()
    yield


app = FastAPI(title="LedgerLink", lifespan=lifespan)


class ClientIn(BaseModel):
    name: str = Field(min_length=1, max_length=200)
    entity_type: EntityType
    financial_year: str = Field(pattern=r"^\d{4}-\d{2}$")
    tb_updated_on: date | None = None
    review_status: str = "Not started"


class ClientOut(ClientIn):
    model_config = ConfigDict(from_attributes=True)
    id: int


class CellChange(BaseModel):
    sheet: str
    cell: str = Field(pattern=r"^[A-Za-z]{1,3}[0-9]{1,7}$")
    value: float | str | bool | None = None
    formula: str | None = Field(default=None, pattern=r"^=")


class CellChanges(BaseModel):
    changes: list[CellChange] = Field(min_length=1, max_length=5000)


@app.get("/api/health")
def health():
    return {"status": "ok"}


@app.get("/api/clients", response_model=list[ClientOut])
def list_clients(session: Session = Depends(get_session)):
    return session.scalars(select(Client).order_by(Client.name)).all()


@app.post("/api/clients", response_model=ClientOut, status_code=201)
def create_client(data: ClientIn, session: Session = Depends(get_session)):
    client = Client(**data.model_dump())
    session.add(client)
    session.commit()
    session.refresh(client)
    return client


@app.delete("/api/clients/{client_id}", status_code=204)
def delete_client(client_id: int, session: Session = Depends(get_session)):
    client = session.get(Client, client_id)
    if client is None:
        raise HTTPException(404, "Client not found")
    session.delete(client)
    session.commit()


@app.post("/api/clients/{client_id}/learn", status_code=202)
def start_learning(client_id: int, file: UploadFile, background: BackgroundTasks,
                   session: Session = Depends(get_session), factory: sessionmaker = Depends(get_session_factory)):
    """LEARN MODE: save the uploaded finished workbook and learn it in the background.
    Returns a job id; follow progress at /api/jobs/{job_id}/events."""
    client = session.get(Client, client_id)
    if client is None:
        raise HTTPException(404, "Client not found")
    filename = Path(file.filename or "").name
    if not filename.lower().endswith((".xlsx", ".xlsm")):
        raise HTTPException(422, "Please upload an Excel workbook (.xlsx).")

    folder = STORAGE_DIR / "clients" / str(client_id)
    folder.mkdir(parents=True, exist_ok=True)
    path = folder / f"{datetime.now():%Y%m%d-%H%M%S}-{filename}"
    path.write_bytes(file.file.read())

    client.review_status = READING
    session.commit()
    job = jobs.create(client.id, client.name, filename)
    background.add_task(_learn_job, job, path, factory)
    return {"job_id": job.id}


def _learn_job(job: jobs.Job, path: Path, factory: sessionmaker):
    """Runs after the upload request has returned."""
    history, dictionary = _ai_context(factory, job.client_id)
    try:
        result = learn_workbook(path, progress=job)
    except Exception as error:
        path.unlink(missing_ok=True)  # nothing learnt, so do not keep the upload
        _set_review_status(factory, job.client_id, NOT_READ)
        if isinstance(error, (ValueError, KeyError, zipfile.BadZipFile, InvalidFileException)):
            job.fail(f"Could not read this workbook. {error}")
        else:
            logger.exception("Learning %s failed", path)
            job.fail("Something went wrong while reading this workbook.")
        return

    with factory() as session:
        client = session.get(Client, job.client_id)
        if client is None:  # deleted while we were reading
            job.fail("This client was deleted while the workbook was being read.")
            return
        mapping, report = result.mapping(), result.report()
        mapping["ai"] = _ai_suggestions(path, mapping, report, history, dictionary)
        workbook = ClientWorkbook(client=client, financial_year=result.financial_year, file_path=str(path),
                                  mapping=mapping, report=report, source="learn", status="learned")
        session.add(workbook)
        client.review_status = _review_text(build_review(path, mapping, mapping, report))
        client.tb_updated_on = date.today()
        session.commit()
        job.finish(workbook.id)


def _set_review_status(factory: sessionmaker, client_id: int, status: str):
    with factory() as session:
        client = session.get(Client, client_id)
        if client:
            client.review_status = status
            session.commit()


@app.get("/api/jobs/{job_id}")
def get_job(job_id: str):
    return _find_job(job_id).snapshot()


@app.get("/api/jobs/{job_id}/events")
async def job_events(job_id: str):
    """Server-sent events: the job's state each time it changes, until it ends."""
    job = _find_job(job_id)

    async def stream():
        sent = -1
        while True:
            state = job.snapshot()
            if state["version"] != sent:
                sent = state["version"]
                yield f"data: {json.dumps(state)}\n\n"
            if state["status"] != "running":
                return
            await asyncio.sleep(0.2)

    return StreamingResponse(stream(), media_type="text/event-stream", headers={"Cache-Control": "no-cache"})


def _find_job(job_id: str) -> jobs.Job:
    job = jobs.get(job_id)
    if job is None:
        raise HTTPException(404, "Job not found. It may have ended before the server restarted.")
    return job


# ---------------------------------------------------------------- workbook screen

class ReviewAction(BaseModel):
    item: str
    action: Literal["approve", "dismiss", "change"]
    ledger: str | None = None  # Change: map to this ledger (a key from "ledgers")
    cell: str | None = None  # Change: add the ledger to this cell
    value: float | str | None = None  # Change: new value for the cell
    formula: str | None = Field(default=None, pattern=r"^=")


@app.get("/api/clients/{client_id}/workbook")
def get_workbook(client_id: int, session: Session = Depends(get_session)):
    """The client's latest workbook as a Univer spreadsheet, with review colours, traces,
    review items and checks."""
    return _workbook_payload(_latest_workbook(session, client_id))


@app.put("/api/clients/{client_id}/workbook/cells")
def save_cell_changes(client_id: int, data: CellChanges, session: Session = Depends(get_session)):
    """Save cells the CA edited. Edited cells turn plum; changing a cell back removes the edit."""
    workbook = _open_workbook(session, client_id)
    changes = [c.model_dump() for c in data.changes]
    try:
        workbook.mapping = apply_edits(workbook.file_path, workbook.mapping, changes, datetime.now())
    except ValueError as error:
        raise HTTPException(422, str(error))
    review = _review(workbook)
    workbook.client.review_status = _review_text(review)
    session.commit()
    cells = edited_cells(workbook.file_path, workbook.mapping, workbook.report, changes, review["overrides"])
    return {"cells": cells, "review": review["items"], "checks": _public_checks(review)}


@app.post("/api/clients/{client_id}/workbook/review")
def act_on_review_item(client_id: int, data: ReviewAction, session: Session = Depends(get_session)):
    """Approve or Change a review item. The decision is saved in the mapping."""
    workbook = _open_workbook(session, client_id)
    item = next((i for i in _review(workbook)["items"] if i["id"] == data.item), None)
    if item is None:
        raise HTTPException(404, "This review item is not open any more.")
    payload = data.model_dump(exclude={"item", "action"}, exclude_none=True)
    try:
        workbook.mapping = apply_action(workbook.file_path, workbook.mapping, item, data.action, payload,
                                        datetime.now())
    except ValueError as error:
        raise HTTPException(422, str(error))
    result = _workbook_payload(workbook)
    workbook.client.review_status = _review_text({"items": result["review"]})
    session.commit()
    return result


@app.post("/api/clients/{client_id}/workbook/finalise")
def finalise_workbook(client_id: int, session: Session = Depends(get_session)):
    """Lock the year. Allowed only when every review item is cleared."""
    workbook = _open_workbook(session, client_id)
    open_items = _review(workbook)["items"]
    if open_items:
        raise HTTPException(409, f"Clear the {len(open_items)} review items before finalising.")
    workbook.status, workbook.finalised_at = "finalised", datetime.now()
    workbook.client.review_status = "Finalised"
    workbook.client.financial_year = workbook.financial_year or workbook.client.financial_year
    session.commit()
    return _workbook_payload(workbook)


@app.get("/api/clients/{client_id}/workbook/download")
def download_workbook(client_id: int, session: Session = Depends(get_session)):
    """The workbook in the client's own format, with the CA's edits and live formulas."""
    workbook = _latest_workbook(session, client_id)
    content = workbook_bytes(workbook.file_path, workbook.mapping.get("edits", {}))
    name = re.sub(r"[^A-Za-z0-9 &._-]+", "", f"{workbook.client.name} FY {workbook.financial_year}").strip()
    return Response(content, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": f'attachment; filename="{name}.xlsx"'})


@app.post("/api/clients/{client_id}/rollforward", status_code=202)
def start_roll_forward(client_id: int, file: UploadFile, background: BackgroundTasks,
                       session: Session = Depends(get_session),
                       factory: sessionmaker = Depends(get_session_factory)):
    """ROLL-FORWARD MODE: upload next year's Tally TB. Builds next year's workbook from the
    latest learned or finalised one, in the background (follow /api/jobs/{job_id}/events)."""
    client = session.get(Client, client_id)
    if client is None:
        raise HTTPException(404, "Client not found")
    base = session.scalars(select(ClientWorkbook).where(ClientWorkbook.client_id == client_id,
                                                        ClientWorkbook.status.in_(("learned", "finalised")))
                           .order_by(ClientWorkbook.id.desc())).first()
    if base is None:
        draft = session.scalars(select(ClientWorkbook).where(ClientWorkbook.client_id == client_id)).first()
        raise HTTPException(409, f"Finalise FY {draft.financial_year} first; next year's TB rolls forward from a "
                                 f"finalised workbook." if draft else "Read a finished workbook for this client first.")
    filename = Path(file.filename or "").name
    if not filename.lower().endswith((".xlsx", ".xlsm")):
        raise HTTPException(422, "Please upload the Tally trial balance as an Excel file (.xlsx).")

    folder = STORAGE_DIR / "clients" / str(client_id)
    folder.mkdir(parents=True, exist_ok=True)
    stamp = f"{datetime.now():%Y%m%d-%H%M%S}"
    tb_path = folder / f"{stamp}-TB-{filename}"
    tb_path.write_bytes(file.file.read())
    out_path = folder / f"{stamp}-roll-forward.xlsx"

    previous_status = client.review_status
    client.review_status = ROLLING
    session.commit()
    job = jobs.create(client.id, client.name, filename, "Rolling forward", jobs.ROLL_FORWARD_STEPS)
    background.add_task(_roll_forward_job, job, base.id, tb_path, out_path, previous_status, factory)
    return {"job_id": job.id}


def _roll_forward_job(job: jobs.Job, base_id: int, tb_path: Path, out_path: Path, previous_status: str,
                      factory: sessionmaker):
    """Runs after the upload request has returned."""
    with factory() as session:
        base = session.get(ClientWorkbook, base_id)
        base_path, base_mapping = base.file_path, base.mapping
    history, dictionary = _ai_context(factory, job.client_id)
    try:
        facts = roll_forward(base_path, base_mapping, tb_path, out_path, progress=job, history=history)
        result = learn_workbook(out_path)
    except Exception as error:
        out_path.unlink(missing_ok=True)
        _set_review_status(factory, job.client_id, previous_status)
        if isinstance(error, (ValueError, KeyError, zipfile.BadZipFile, InvalidFileException)):
            job.fail(f"Could not roll forward with this trial balance. {error}")
        else:
            logger.exception("Roll-forward with %s failed", tb_path)
            job.fail("Something went wrong while rolling forward.")
        return

    # Decisions to leave ledgers unused carry into the new year (CLAUDE.md: CA corrections repeat).
    mapping = {**result.mapping(), "rollforward": facts, "decisions": {},
               "ignored_ledgers": base_mapping.get("ignored_ledgers", [])}
    report = result.report()
    mapping["ai"] = _ai_suggestions(out_path, mapping, report, history, dictionary)
    review = build_review(out_path, mapping, mapping, report)
    job.step_done(3, "; ".join(c["label"] for c in review["checks"]) + ". " + mapping["ai"]["about"])
    with factory() as session:
        client = session.get(Client, job.client_id)
        if client is None:
            job.fail("This client was deleted while rolling forward.")
            return
        for old_draft in [w for w in client.workbooks if w.status == "draft"]:
            session.delete(old_draft)  # a new TB for the same year replaces the unfinished draft
        workbook = ClientWorkbook(client=client, financial_year=result.financial_year, file_path=str(out_path),
                                  mapping=mapping, report=report, source="roll-forward", status="draft")
        session.add(workbook)
        client.financial_year = result.financial_year
        client.tb_updated_on = date.today()
        client.review_status = _review_text(review)
        session.commit()
        job.finish(workbook.id)


@app.post("/api/clients/{client_id}/build", status_code=202)
def start_build(client_id: int, file: UploadFile, background: BackgroundTasks, sheets: list[str] = Form(...),
                session: Session = Depends(get_session), factory: sessionmaker = Depends(get_session_factory)):
    """BUILD MODE: statements from a Tally TB only, in the default format for the entity type.
    Only the ticked sheets are made. Runs in the background (follow /api/jobs/{job_id}/events)."""
    client = session.get(Client, client_id)
    if client is None:
        raise HTTPException(404, "Client not found")
    unknown = [s for s in sheets if s not in BUILD_SHEETS]
    if unknown or not sheets:
        raise HTTPException(422, f"Choose sheets from: {', '.join(BUILD_SHEETS)}.")
    filename = Path(file.filename or "").name
    if not filename.lower().endswith((".xlsx", ".xlsm")):
        raise HTTPException(422, "Please upload the Tally trial balance as an Excel file (.xlsx).")

    folder = STORAGE_DIR / "clients" / str(client_id)
    folder.mkdir(parents=True, exist_ok=True)
    stamp = f"{datetime.now():%Y%m%d-%H%M%S}"
    tb_path = folder / f"{stamp}-TB-{filename}"
    tb_path.write_bytes(file.file.read())
    out_path = folder / f"{stamp}-built.xlsx"

    client.review_status = BUILDING
    session.commit()
    job = jobs.create(client.id, client.name, filename, "Building", jobs.BUILD_STEPS)
    background.add_task(_build_job, job, client.entity_type, sheets, tb_path, out_path, factory)
    return {"job_id": job.id}


def _build_job(job: jobs.Job, entity_type: str, sheets: list[str], tb_path: Path, out_path: Path,
               factory: sessionmaker):
    """Runs after the upload request has returned."""
    history, dictionary = _ai_context(factory, job.client_id)
    try:
        facts = build_workbook(tb_path, entity_type, job.client_name, sheets, out_path, progress=job)
        result = learn_workbook(out_path)
    except Exception as error:
        out_path.unlink(missing_ok=True)
        _set_review_status(factory, job.client_id, NOT_READ)
        if isinstance(error, (ValueError, KeyError, zipfile.BadZipFile, InvalidFileException)):
            job.fail(f"Could not build from this trial balance. {error}")
        else:
            logger.exception("Build from %s failed", tb_path)
            job.fail("Something went wrong while building the statements.")
        return

    mapping, report = {**result.mapping(), "build": facts, "decisions": {}}, result.report()
    mapping["ai"] = _ai_suggestions(out_path, mapping, report, history, dictionary)
    review = build_review(out_path, mapping, mapping, report)
    job.step_done(3, "; ".join(c["label"] for c in review["checks"]) + ". " + mapping["ai"]["about"])
    with factory() as session:
        client = session.get(Client, job.client_id)
        if client is None:
            job.fail("This client was deleted while building.")
            return
        workbook = ClientWorkbook(client=client, financial_year=result.financial_year, file_path=str(out_path),
                                  mapping=mapping, report=report, source="build", status="draft")
        session.add(workbook)
        client.financial_year = result.financial_year or client.financial_year
        client.tb_updated_on = date.today()
        client.review_status = _review_text(review)
        session.commit()
        job.finish(workbook.id)


def _ai_context(factory: sessionmaker, client_id: int) -> tuple[dict, dict]:
    """The AI layer's evidence: this client's own earlier workbooks (1st), and the firm-wide
    dictionary from every other client's learned or finalised workbook (2nd)."""
    with factory() as session:
        own = session.scalars(select(ClientWorkbook).where(ClientWorkbook.client_id == client_id)).all()
        others = session.scalars(select(ClientWorkbook).where(
            ClientWorkbook.client_id != client_id, ClientWorkbook.status.in_(("learned", "finalised")))).all()
        return ai.client_history([w.mapping for w in own]), ai.firm_dictionary([w.mapping for w in others])


def _ai_suggestions(path: Path, mapping: dict, report: dict, history: dict, dictionary: dict) -> dict:
    """AI suggestions for ledgers no formula uses; each shows amber until the CA approves it.
    Cells that already have an exact "link to TB" review item are left to that item."""
    if mapping.get("rollforward"):
        links = mapping["rollforward"]["suggestions"]
    else:
        links = link_suggestions(report["manual_cells"], parse_trial_balance(path))
    covered = {s["cell"] for s in links}
    return {"targets": ai.suggest_for_unused(mapping, report, history, dictionary, covered), "about": ai.describe()}


def _review(workbook: ClientWorkbook) -> dict:
    """Review items, checks and colours for the workbook as the CA has left it."""
    edits = workbook.mapping.get("edits", {})
    if edits:
        live = learn_effective(workbook.file_path, edits)
        live_mapping, live_report = live.mapping(), live.report()
    else:
        live_mapping, live_report = workbook.mapping, workbook.report
    return build_review(workbook.file_path, workbook.mapping, live_mapping, live_report)


def _review_text(review: dict) -> str:
    count = len(review["items"])
    return f"{count} to check" if count else "Ready to finalise"


def _public_checks(review: dict) -> list[dict]:
    return [{"id": c["id"], "ok": c["ok"], "label": c["label"], "detail": c["detail"]} for c in review["checks"]]


def _workbook_payload(workbook: ClientWorkbook) -> dict:
    review = _review(workbook)
    view = workbook_view(workbook.file_path, workbook.mapping, workbook.report, review["overrides"])
    tb = parse_trial_balance(workbook.file_path)
    fy = workbook.financial_year
    # A roll-forward draft can take a corrected TB for the same year; otherwise it is next year's.
    redo = workbook.status == "draft" and workbook.source == "roll-forward"
    next_fy = fy if redo or not fy else f"{int(fy[:4]) + 1}-{(int(fy[:4]) + 2) % 100:02d}"
    return {
        "workbook_id": workbook.id, "financial_year": fy, "status": workbook.status, "source": workbook.source,
        "finalised_at": workbook.finalised_at.isoformat(timespec="seconds") if workbook.finalised_at else None,
        "next_financial_year": next_fy, **view,
        "review": review["items"], "checks": _public_checks(review),
        "ledgers": [{"key": r.key, "name": r.name, "group": r.group_path} for r in tb.rows],
    }


def _open_workbook(session: Session, client_id: int) -> ClientWorkbook:
    """The latest workbook, refusing changes once it is finalised."""
    workbook = _latest_workbook(session, client_id)
    if workbook.status == "finalised":
        raise HTTPException(409, f"FY {workbook.financial_year} is finalised and locked.")
    return workbook


def _latest_workbook(session: Session, client_id: int) -> ClientWorkbook:
    if session.get(Client, client_id) is None:
        raise HTTPException(404, "Client not found")
    workbook = session.scalars(select(ClientWorkbook).where(ClientWorkbook.client_id == client_id)
                               .order_by(ClientWorkbook.id.desc())).first()
    if workbook is None:
        raise HTTPException(404, "No workbook has been read for this client yet.")
    return workbook
