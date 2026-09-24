"""What the Workbook screen shows: the client's workbook as a Univer spreadsheet, the review
colour of each cell and a plain-words trace of where each figure comes from.

The original upload is never changed. CA edits are kept in the mapping under "edits"
(CLAUDE.md: CA corrections are saved to the mapping) and laid over the file here.
"""
import re
from datetime import date, datetime, time
from pathlib import Path

import openpyxl
from openpyxl.cell.cell import MergedCell
from openpyxl.utils import range_boundaries

from app.learn import INTERNAL, MANUAL, PRIOR_YEAR, TB_LINKED
from app.tally_tb import format_inr, parse_trial_balance

# Review colours; same values as the tokens in frontend/src/styles.css.
GREEN, AMBER, RED, PLUM = "g", "a", "r", "p"  # linked as last year, AI suggestion, needs CA, edited by CA
STATUS_BG = {GREEN: "#E7F2EA", AMBER: "#FAEFD6", RED: "#FAE3E0", PLUM: "#EEE6F5"}
STATUS_INK = {GREEN: "#1C6B3A", AMBER: "#7A5200", RED: "#A1261B", PLUM: "#5B2C83"}

# Indian grouping (1,36,74,167.85). Excel formats allow two conditions, so this is exact for
# positive amounts below 100 crore; larger or negative amounts fall back to #,##0 grouping.
INDIAN_FORMAT = '[>=10000000]#","##","##","##0{d};[>=100000]#","##","##0{d};#,##0{d}'
PLAIN_NUMBER = re.compile(r'(?:"")?0(?:\.0+)?')  # 0, 0.00, ""0.00
EXCEL_EPOCH = datetime(1899, 12, 30)
BORDER_STYLES = {"thin": 1, "hair": 2, "dotted": 3, "dashed": 4, "dashDot": 5, "dashDotDot": 6, "double": 7,
                 "medium": 8, "mediumDashed": 9, "mediumDashDot": 10, "mediumDashDotDot": 11,
                 "slantDashDot": 12, "thick": 13}
H_ALIGN = {"left": 1, "center": 2, "centerContinuous": 2, "right": 3, "justify": 4}
V_ALIGN = {"top": 1, "center": 2, "bottom": 3}
EXTRA_ROWS, EXTRA_COLS = 50, 10  # empty space below and right of the used area, as in Excel


def workbook_view(path: str | Path, mapping: dict, report: dict, overrides: dict | None = None) -> dict:
    """Everything the Workbook screen needs: the Univer snapshot plus status and trace per cell.
    overrides: {cell: {"status", "reason"}} from the review (roll-forward matches, CA decisions)."""
    cells = cell_details(path, mapping, report, overrides)
    snapshot = univer_snapshot(path, mapping.get("edits", {}), cells)
    return {"snapshot": snapshot, "cells": cells}


# ---------------------------------------------------------------- statuses and traces

def cell_details(path: str | Path, mapping: dict, report: dict,
                 overrides: dict | None = None) -> dict[str, dict[str, dict]]:
    """{sheet: {"D21": {"status": "r" | None, "trace": [what, from], "reason": why}}} for every
    cell worth explaining."""
    details: dict[str, dict[str, dict]] = {}
    cells = mapping["cells"]
    precedents = mapping.get("precedents", {})
    ledgers = mapping["ledgers"]
    issues = {i["cell"] for i in report["inconsistencies"]}
    typed_instead = {m["cell"]: m["matches_ledgers"] for m in report["manual_cells"] if m["matches_ledgers"]}

    for key, cell in cells.items():
        status = None
        if key in issues or key in typed_instead:
            status = RED  # inconsistent with last year, or TB amount typed in instead of linked
        elif cell["category"] in (TB_LINKED, INTERNAL):
            status = GREEN
        what = cell["label"] or key
        source = _source(key, cell, cells, ledgers, typed_instead, precedents.get(key, []))
        _put(details, key, status, [what, source])

    # TB sheet: every row says where it is used; unused balances are red. Rows are read from
    # the file itself because the mapping is keyed by ledger name, not row.
    unused = {u["ledger"] for u in report["unused_ledgers"]}
    tb = parse_trial_balance(path)
    for row in tb.rows:
        used_in = sorted({u["cell"] for u in ledgers.get(row.key, {}).get("used_in", [])})
        source = ("used in " + _list(used_in)) if used_in else "not used anywhere in the statements"
        for col in (tb.name_column, tb.debit_column, tb.credit_column):
            _put(details, f"{_quote(tb.sheet_name)}!{col}{row.row}", RED if row.key in unused else None,
                 [row.name, source])

    for key, override in (overrides or {}).items():
        sheet, address = _split_key(key)
        known = details.get(sheet, {}).get(address)
        _put(details, key, override["status"], known["trace"] if known else [address, ""], override["reason"])

    for sheet, sheet_edits in mapping.get("edits", {}).items():
        for address, edit in sheet_edits.items():
            _put(details, f"{_quote(sheet)}!{address}", PLUM, [_edited_label(details, sheet, address), _edit_text(edit)])
    return details


def _source(key: str, cell: dict, cells: dict, ledgers: dict, typed_instead: dict, reads: list[str]) -> str:
    """Plain words for where a figure comes from."""
    category, formula = cell["category"], cell["formula"]
    if category == TB_LINKED:
        names = list(dict.fromkeys(ledgers[link["ledger"]]["name"] for link in cell["links"]))
        return "TB: " + _list(names)
    if category == MANUAL:
        typed = f"typed in ({formula.lstrip('=')})" if formula else "typed in"
        if key in typed_instead:
            return f"{typed}; equals TB {_list(typed_instead[key])}, not linked"
        return typed
    if category == PRIOR_YEAR and not formula:
        return "last year's figure, typed in"
    if category == INTERNAL:
        return _list([_describe(ref, cells) for ref in reads]) or formula
    return formula or ""


def _describe(ref: str, cells: dict) -> str:
    """"'BS Schedules'!C58" -> "BS Schedules: Total Cash & Bank"."""
    sheet = ref.rsplit("!", 1)[0].strip("'").replace("''", "'")
    label = cells.get(ref, {}).get("label")
    return f"{sheet}: {label}" if label else ref


def _put(details: dict, key: str, status: str | None, trace: list[str], reason: str | None = None):
    sheet, address = _split_key(key)
    details.setdefault(sheet, {})[address] = {"status": status, "trace": trace, "reason": reason}


def _split_key(key: str) -> tuple[str, str]:
    sheet, address = key.rsplit("!", 1)
    return (sheet[1:-1].replace("''", "'") if sheet.startswith("'") else sheet), address


def _edited_label(details: dict, sheet: str, address: str) -> str:
    known = details.get(sheet, {}).get(address)
    return known["trace"][0] if known else address


def _edit_text(edit: dict) -> str:
    when = datetime.fromisoformat(edit["edited_at"])
    value = edit.get("value")
    if edit.get("formula"):
        content = edit["formula"]
    elif value is None or value == "":
        content = "cleared"
    elif isinstance(value, (int, float)) and not isinstance(value, bool):
        content = format_inr(value)
    else:
        content = str(value)
    return f"edited by CA on {when.day} {when:%b %Y}: {content}"


def _list(items: list[str], limit: int = 4) -> str:
    if len(items) <= limit:
        return ", ".join(items)
    return ", ".join(items[:limit]) + f" and {len(items) - limit} more"


def _quote(sheet: str) -> str:
    return sheet if sheet.replace("_", "").isalnum() else "'" + sheet.replace("'", "''") + "'"


# ---------------------------------------------------------------- Univer snapshot

def univer_snapshot(path: str | Path, edits: dict, details: dict) -> dict:
    """The workbook in Univer's IWorkbookData format: values, formulas, styles, merges, sizes."""
    wb = openpyxl.load_workbook(path)
    saved = openpyxl.load_workbook(path, data_only=True)
    styles = _StyleTable()
    sheets, order = {}, []
    for index, ws in enumerate(wb.worksheets):
        sheet_id = f"sheet{index}"
        order.append(sheet_id)
        sheets[sheet_id] = _sheet(ws, saved[ws.title], edits.get(ws.title, {}), details.get(ws.title, {}), styles)
        sheets[sheet_id]["id"] = sheet_id
    return {"id": "workbook", "name": Path(path).stem, "locale": "enUS", "styles": styles.by_id,
            "sheetOrder": order, "sheets": sheets}


def _status_style(status: str) -> dict:
    """Review colour as a cell style: tinted fill and matching text colour."""
    return {"bg": {"rgb": STATUS_BG[status]}, "cl": {"rgb": STATUS_INK[status]}}


class _StyleTable:
    """Univer keeps each style once per workbook; cells refer to it by id."""

    def __init__(self):
        self.by_id: dict[str, dict] = {}
        self._ids: dict[str, str] = {}

    def id_for(self, style: dict) -> str | None:
        if not style:
            return None
        text = repr(sorted(style.items()))
        if text not in self._ids:
            self._ids[text] = f"s{len(self._ids)}"
            self.by_id[self._ids[text]] = style
        return self._ids[text]


def _sheet(ws, saved_ws, edits: dict, details: dict, styles: _StyleTable) -> dict:
    cell_data: dict[int, dict[int, dict]] = {}
    for row in ws.iter_rows():
        for c in row:
            if isinstance(c, MergedCell) or (c.value is None and not c.has_style):
                continue
            cell = _cell(c.value, saved_ws.cell(c.row, c.column).value, c.number_format)
            style = _style(c)
            status = details.get(c.coordinate, {}).get("status")
            if status:
                style.update(_status_style(status))
            if (sid := styles.id_for(style)) is not None:
                cell["s"] = sid
            if cell:
                cell_data.setdefault(c.row - 1, {})[c.column - 1] = cell

    for address, edit in edits.items():  # CA edits laid over the original file
        col, row, _, _ = range_boundaries(address)
        cell = cell_data.setdefault(row - 1, {}).setdefault(col - 1, {})
        for field in ("v", "f", "t"):
            cell.pop(field, None)
        cell.update(_cell(edit.get("formula") or edit.get("value"), None, "General"))
        cell["s"] = styles.id_for({**styles.by_id.get(cell.get("s"), {}), **_status_style(PLUM)})

    max_col = max(ws.max_column, max((c + 1 for r in cell_data.values() for c in r), default=1))
    max_row = max(ws.max_row, max((r + 1 for r in cell_data), default=1))
    column_count = max_col + EXTRA_COLS
    return {
        "name": ws.title,
        "rowCount": max_row + EXTRA_ROWS,
        "columnCount": column_count,
        "cellData": cell_data,
        "mergeData": [_merge(m) for m in ws.merged_cells.ranges],
        "columnData": {i - 1: {"w": _col_px(d.width)} for i, d in _column_widths(ws, column_count).items()},
        "rowData": {r - 1: {"h": round(d.height * 4 / 3), "hd": int(bool(d.hidden))}
                    for r, d in ws.row_dimensions.items() if d.height},
        "showGridlines": 0 if ws.sheet_view.showGridLines is False else 1,
        "hidden": int(ws.sheet_state != "visible"),
        "defaultColumnWidth": 72,
        "defaultRowHeight": 20,
    }


def _cell(value, saved_value, number_format: str) -> dict:
    formula = value if isinstance(value, str) and value.startswith("=") and len(value) > 1 else None
    formula = formula or (value.text if hasattr(value, "text") and isinstance(value.text, str) else None)
    if formula:
        cell = {"f": formula}
        if saved_value is not None:
            cell.update(_typed(saved_value))
        return cell
    return _typed(value) if value is not None else {}


def _typed(value) -> dict:
    """Univer cell value and type: 1 text, 2 number, 3 boolean."""
    if isinstance(value, bool):
        return {"v": int(value), "t": 3}
    if isinstance(value, (int, float)):
        return {"v": value, "t": 2}
    if isinstance(value, (datetime, date)):
        moment = value if isinstance(value, datetime) else datetime.combine(value, time())
        delta = moment - EXCEL_EPOCH
        return {"v": delta.days + delta.seconds / 86400, "t": 2}
    if isinstance(value, time):
        return {"v": (value.hour * 3600 + value.minute * 60 + value.second) / 86400, "t": 2}
    return {"v": str(value), "t": 1}


def _style(c) -> dict:
    """openpyxl formatting -> Univer IStyleData."""
    style: dict = {}
    font = c.font
    if font.name:
        style["ff"] = font.name
    if font.sz:
        style["fs"] = float(font.sz)
    if font.b:
        style["bl"] = 1
    if font.i:
        style["it"] = 1
    if font.u and font.u != "none":
        style["ul"] = {"s": 1}
    if font.strike:
        style["st"] = {"s": 1}
    if (colour := _rgb(font.color)):
        style["cl"] = {"rgb": colour}
    if c.fill is not None and c.fill.fill_type == "solid" and (colour := _rgb(c.fill.fgColor)):
        style["bg"] = {"rgb": colour}

    align = c.alignment
    if align.horizontal in H_ALIGN:
        style["ht"] = H_ALIGN[align.horizontal]
    if align.vertical in V_ALIGN:
        style["vt"] = V_ALIGN[align.vertical]
    if align.wrap_text:
        style["tb"] = 3  # wrap
    if align.indent:
        style["pd"] = {"l": int(align.indent) * 9}

    borders = {}
    for side, key in (("top", "t"), ("bottom", "b"), ("left", "l"), ("right", "r")):
        edge = getattr(c.border, side)
        if edge is not None and edge.style in BORDER_STYLES:
            borders[key] = {"s": BORDER_STYLES[edge.style], "cl": {"rgb": _rgb(edge.color) or "#000000"}}
    if borders:
        style["bd"] = borders

    if (pattern := _number_format(c.number_format)):
        style["n"] = {"pattern": pattern}
    return style


def _number_format(excel_format: str | None) -> str | None:
    """Keep the workbook's format, but show grouped amounts the Indian way."""
    if not excel_format or excel_format == "General":
        return None
    first = excel_format.split(";")[0]
    # Grouped amounts (#,##0.00, accounting formats) and plain ones (0.00, as Tally exports)
    if "#,##0" in excel_format or PLAIN_NUMBER.fullmatch(first):
        decimals = len(m.group(1)) if (m := re.search(r"0\.(0+)", first)) else 0
        return INDIAN_FORMAT.format(d="." + "0" * decimals if decimals else "")
    return excel_format


def _rgb(colour) -> str | None:
    """openpyxl colour -> "#RRGGBB". Theme and indexed colours fall back to Univer's defaults."""
    if colour is None or getattr(colour, "type", None) != "rgb" or not isinstance(colour.rgb, str):
        return None
    value = colour.rgb[-6:]
    return f"#{value.upper()}" if re.fullmatch(r"[0-9A-Fa-f]{6}", value) else None


def _merge(merged) -> dict:
    return {"startRow": merged.min_row - 1, "endRow": merged.max_row - 1,
            "startColumn": merged.min_col - 1, "endColumn": merged.max_col - 1}


def _column_widths(ws, column_count: int) -> dict[int, object]:
    """openpyxl groups columns (min..max share one width, sometimes up to column XFD);
    spread them per column, only as far as the sheet is shown."""
    widths = {}
    for dim in ws.column_dimensions.values():
        if dim.width:
            for index in range(dim.min or 1, min(dim.max or dim.min or 1, column_count) + 1):
                widths[index] = dim
    return widths


def _col_px(width_chars: float) -> int:
    return round(width_chars * 7 + 5)  # Excel character width to pixels (Calibri 11)


# ---------------------------------------------------------------- CA edits

def apply_edits(path: str | Path, mapping: dict, changes: list[dict], now: datetime) -> dict:
    """Return a new mapping with the CA's cell changes saved under "edits".

    Each change is {"sheet", "cell", "value", "formula"}. A change that puts a cell back to
    what the uploaded file has (e.g. undo) removes the edit instead.
    """
    wb = openpyxl.load_workbook(path)
    edits = {sheet: dict(cells) for sheet, cells in mapping.get("edits", {}).items()}
    for change in changes:
        sheet, address = change["sheet"], change["cell"].upper()
        if sheet not in wb.sheetnames:
            raise ValueError(f"No sheet named {sheet!r}.")
        range_boundaries(address)  # raises ValueError for a bad address
        formula, value = change.get("formula") or None, change.get("value")
        if _same_as_file(wb[sheet][address].value, formula, value):
            edits.get(sheet, {}).pop(address, None)
        else:
            edits.setdefault(sheet, {})[address] = {"formula": formula, "value": None if formula else value,
                                                    "edited_at": now.isoformat(timespec="seconds")}
    return {**mapping, "edits": {sheet: cells for sheet, cells in edits.items() if cells}}


def _same_as_file(original, formula: str | None, value) -> bool:
    if formula:
        return isinstance(original, str) and original.replace(" ", "") == formula.replace(" ", "")
    if original is None or original == "":
        return value is None or value == ""
    if isinstance(original, (int, float)) and not isinstance(original, bool):
        try:
            return abs(float(value) - original) < 1e-9
        except (TypeError, ValueError):
            return False
    return str(original) == str(value)


def edited_cells(path: str | Path, mapping: dict, report: dict, changes: list[dict],
                 overrides: dict | None = None) -> list[dict]:
    """After saving: status, trace and the colours to show now for each changed cell: the review
    colour, else the cell's own fill and font colour in the file (e.g. after an undo)."""
    details = cell_details(path, mapping, report, overrides)
    wb = openpyxl.load_workbook(path)
    result = []
    for change in changes:
        sheet, address = change["sheet"], change["cell"].upper()
        info = details.get(sheet, {}).get(address, {"status": None, "trace": None, "reason": None})
        own = wb[sheet][address]
        if info["status"]:
            bg, ink = STATUS_BG[info["status"]], STATUS_INK[info["status"]]
        else:
            bg = _rgb(own.fill.fgColor) if own.fill is not None and own.fill.fill_type == "solid" else None
            ink = _rgb(own.font.color)
        result.append({"sheet": sheet, "cell": address, **info, "bg": bg, "ink": ink})
    return result
