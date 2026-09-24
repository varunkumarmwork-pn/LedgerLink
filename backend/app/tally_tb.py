"""Read a Tally trial balance (TB) export from an Excel workbook.

Tally marks the row type through formatting (see CLAUDE.md, "Tally TB export"):
  GROUP      bold, indent 0            e.g. "Current Assets"     (carries a subtotal)
  SUB-GROUP  not italic, indent > 0    e.g. "Bank Accounts"      (carries a subtotal)
  LEDGER     italic                    e.g. "Bank of Baroda"
Debit is in one column and Credit in the next. The last row is "Grand Total".

Indentation alone does not always say which group a row sits under: in the Venus file the
ledger "G S T" has the same indent as its sub-group "Duties & Taxes", while "Inventory
Difference Account" has the same indent as "Bank Accounts" but belongs to "Current Assets".
So we also use the subtotals: a group keeps taking rows only while its subtotal is not yet
used up, and only if the next row fits inside what is left.
"""
import re
from dataclasses import dataclass, field
from decimal import Decimal, InvalidOperation
from pathlib import Path

import openpyxl
from openpyxl.utils import get_column_letter
from openpyxl.workbook import Workbook
from openpyxl.worksheet.worksheet import Worksheet

GROUP, SUB_GROUP, LEDGER = "group", "sub-group", "ledger"
ZERO = Decimal("0.00")
GSTIN_PATTERN = re.compile(r"\d{2}[A-Z]{5}\d{4}[A-Z][A-Z\d]Z[A-Z\d]")
PERIOD_PATTERN = re.compile(r"\d{1,2}-[A-Za-z]{3}-\d{2,4}\s+to\s+\d{1,2}-[A-Za-z]{3}-\d{2,4}")
TB_SHEET_NAME = re.compile(r"^\s*(tb|trial\s*balance)\b", re.IGNORECASE)
HEADER_SEARCH_ROWS = 40  # the "Particulars" header is always near the top
# Tally prints Closing Stock in the TB for information but leaves it out of the Grand Total
# (it is both an asset and a P&L credit, so it would not balance otherwise).
MEMO_LINES = {"closing stock"}


@dataclass
class TBRow:
    row: int  # Excel row number in the TB sheet
    name: str
    kind: str  # GROUP, SUB_GROUP or LEDGER
    indent: int
    dr: Decimal
    cr: Decimal
    path: tuple[str, ...] = ()  # names of the groups above this row, top first
    memo: bool = False  # shown in the TB but not part of the Grand Total (Closing Stock)
    children: list["TBRow"] = field(default_factory=list, repr=False)

    @property
    def group_path(self) -> str:
        return " > ".join(self.path)

    @property
    def key(self) -> str:
        """Stable name for this row across years, e.g. "Current Assets > Bank Accounts > Bank of Baroda".
        Mappings are stored by this key, never by row number."""
        return " > ".join(self.path + (self.name,))

    def with_descendants(self) -> list["TBRow"]:
        return [self] + [d for child in self.children for d in child.with_descendants()]


@dataclass
class TrialBalance:
    sheet_name: str
    company: str
    address: list[str]
    gstin: str | None
    period: str | None
    name_column: str  # column letter of the ledger names, e.g. "A"
    debit_column: str  # column letter, e.g. "B"
    credit_column: str
    rows: list[TBRow]
    total_dr: Decimal  # as printed on the "Grand Total" row
    total_cr: Decimal
    problems: list[str]  # plain-words issues found while checking the TB

    @property
    def ledgers(self) -> list[TBRow]:
        return [r for r in self.rows if r.kind == LEDGER]

    @property
    def balances(self) -> list[TBRow]:
        """Every row that holds a balance with nothing beneath it.

        That is all ledgers plus groups Tally shows collapsed (e.g. "Sundry Debtors",
        "Closing Stock"). Adding up the non-memo ones gives the TB totals without double counting.
        """
        return [r for r in self.rows if not r.children and (r.dr or r.cr)]

    @property
    def is_balanced(self) -> bool:
        return self.total_dr == self.total_cr


def parse_trial_balance(source: str | Path | Workbook) -> TrialBalance:
    """Find the TB sheet in a workbook (path or open workbook) and read it."""
    wb = source if isinstance(source, Workbook) else openpyxl.load_workbook(source, data_only=True)
    ws, header = find_tb_sheet(wb)
    name_col, debit_col, first_data_row = header

    rows, grand_total = _read_rows(ws, name_col, debit_col, first_data_row)
    problems = _build_tree(rows)
    if grand_total is None:
        problems.append("No 'Grand Total' row found at the end of the trial balance.")
        grand_total = (ZERO, ZERO)
    total_dr, total_cr = grand_total
    problems += _check_totals(rows, total_dr, total_cr)

    company, address, gstin, period = _read_company_details(ws, name_col, first_data_row)
    return TrialBalance(
        sheet_name=ws.title,
        company=company,
        address=address,
        gstin=gstin,
        period=period,
        name_column=get_column_letter(name_col),
        debit_column=get_column_letter(debit_col),
        credit_column=get_column_letter(debit_col + 1),
        rows=rows,
        total_dr=total_dr,
        total_cr=total_cr,
        problems=problems,
    )


# ---------------------------------------------------------------- finding the sheet

def find_tb_sheet(wb: Workbook) -> tuple[Worksheet, tuple[int, int, int]]:
    """Return the TB sheet and its header (name column, debit column, first data row).

    Sheets named like "TB", "Trial Balance" or "TB 25-26" are tried first, then every
    other sheet is checked by content.
    """
    by_name = [ws for ws in wb.worksheets if TB_SHEET_NAME.match(ws.title)]
    others = [ws for ws in wb.worksheets if ws not in by_name]
    for ws in by_name + others:
        header = _find_header(ws)
        if header:
            return ws, header
    raise ValueError("No Tally trial balance found: no sheet has a 'Particulars' header with Debit and Credit columns.")


def _find_header(ws: Worksheet) -> tuple[int, int, int] | None:
    """Find "Particulars" and, a few rows below it, a "Debit" cell followed by "Credit"."""
    for row in ws.iter_rows(max_row=min(ws.max_row, HEADER_SEARCH_ROWS)):
        for cell in row:
            if _text(cell.value).lower() != "particulars":
                continue
            for r in range(cell.row, cell.row + 6):
                debit_cols = [c.column for c in ws[r] if _text(c.value).lower() == "debit"]
                # Tally puts the closing balance last, so take the last Debit column.
                for col in reversed(debit_cols):
                    if _text(ws.cell(r, col + 1).value).lower() == "credit":
                        return cell.column, col, r + 1
    return None


# ---------------------------------------------------------------- reading rows

def _read_rows(ws: Worksheet, name_col: int, debit_col: int, first_row: int):
    """Read every TB row down to "Grand Total". Returns (rows, (grand_dr, grand_cr) or None)."""
    rows = []
    for r in range(first_row, ws.max_row + 1):
        name_cell = ws.cell(r, name_col)
        name = _text(name_cell.value)
        if not name:
            continue
        dr, cr = _amount(ws.cell(r, debit_col).value), _amount(ws.cell(r, debit_col + 1).value)
        if name.lower() == "grand total":
            return rows, (dr, cr)
        indent = int(name_cell.alignment.indent or 0)
        rows.append(TBRow(row=r, name=name, kind=_row_kind(name_cell.font, indent), indent=indent, dr=dr, cr=cr,
                          memo=name.lower() in MEMO_LINES))
    return rows, None


def _row_kind(font, indent: int) -> str:
    if font.i:
        return LEDGER
    # Indent 0 is a Tally primary group. This also covers "Profit & Loss A/c", which Tally
    # prints at indent 0 without bold.
    return GROUP if indent == 0 else SUB_GROUP


def _build_tree(rows: list[TBRow]) -> list[str]:
    """Attach each row to its parent group; set .path and .children. Returns problems."""
    problems = []
    open_groups: list[list] = []  # stack of [group row, Dr still to place, Cr still to place]

    for row in rows:
        if row.kind == GROUP:
            open_groups = [[row, row.dr, row.cr]]
            continue

        while open_groups and not _belongs_to(row, *open_groups[-1]):
            open_groups.pop()
        if not open_groups:
            problems.append(f"Row {row.row} '{row.name}' does not fit under any group subtotal.")
            continue

        parent = open_groups[-1]
        parent[0].children.append(row)
        parent[1] -= row.dr
        parent[2] -= row.cr
        row.path = parent[0].path + (parent[0].name,)
        if row.kind == SUB_GROUP:
            open_groups.append([row, row.dr, row.cr])
    return problems


def _belongs_to(row: TBRow, group: TBRow, dr_left: Decimal, cr_left: Decimal) -> bool:
    if dr_left == ZERO and cr_left == ZERO:
        return False  # this group's subtotal is fully used up
    if row.kind == SUB_GROUP and group.kind == SUB_GROUP and row.indent <= group.indent:
        return False  # a sibling sub-group, not a child
    return row.dr <= dr_left and row.cr <= cr_left


# ---------------------------------------------------------------- checks

def _check_totals(rows: list[TBRow], total_dr: Decimal, total_cr: Decimal) -> list[str]:
    problems = []
    if total_dr != total_cr:
        problems.append(
            f"Debit total {format_inr(total_dr)} does not equal credit total {format_inr(total_cr)} "
            f"(difference {format_inr(total_dr - total_cr)})."
        )

    groups = [r for r in rows if r.kind == GROUP and not r.memo]
    group_dr, group_cr = sum((g.dr for g in groups), ZERO), sum((g.cr for g in groups), ZERO)
    if (group_dr, group_cr) != (total_dr, total_cr):
        problems.append(
            f"Groups add up to Dr {format_inr(group_dr)} / Cr {format_inr(group_cr)}, "
            f"but Grand Total shows Dr {format_inr(total_dr)} / Cr {format_inr(total_cr)}."
        )

    for g in rows:
        if not g.children:
            continue
        child_dr, child_cr = sum((c.dr for c in g.children), ZERO), sum((c.cr for c in g.children), ZERO)
        if (child_dr, child_cr) != (g.dr, g.cr):
            problems.append(
                f"'{g.name}' (row {g.row}) shows Dr {format_inr(g.dr)} / Cr {format_inr(g.cr)}, "
                f"but its rows add up to Dr {format_inr(child_dr)} / Cr {format_inr(child_cr)}."
            )
    return problems


# ---------------------------------------------------------------- company details

def _read_company_details(ws: Worksheet, name_col: int, first_data_row: int):
    """Company name, address lines, GSTIN and period from the lines above the header."""
    lines = []
    for r in range(1, first_data_row):
        text = _text(ws.cell(r, name_col).value)
        if text.lower() == "particulars":
            break
        if text:
            lines.append(text)

    company = lines[0] if lines else ""
    address, gstin, period = [], None, None
    for line in lines[1:]:
        gst_match = GSTIN_PATTERN.search(line.upper())
        period_match = PERIOD_PATTERN.search(line)
        if gst_match:
            gstin = gst_match.group(0)
        elif period_match:
            period = period_match.group(0)
        elif line.lower() != "trial balance":
            address.append(line)
    return company, address, gstin, period


# ---------------------------------------------------------------- helpers

def _text(value) -> str:
    return "" if value is None else str(value).strip()


def _amount(value) -> Decimal:
    """Excel cell value to rupees and paise. Blank or text becomes 0."""
    if value is None or value == "":
        return ZERO
    try:
        return Decimal(str(value).replace(",", "")).quantize(Decimal("0.01"))
    except InvalidOperation:
        return ZERO


def format_inr(amount: Decimal | float) -> str:
    """Indian number format: 101252474.34 -> 10,12,52,474.34"""
    sign = "-" if amount < 0 else ""
    rupees, paise = f"{abs(amount):.2f}".split(".")
    if len(rupees) > 3:
        head, last3 = rupees[:-3], rupees[-3:]
        pairs = [head[max(i - 2, 0):i] for i in range(len(head), 0, -2)][::-1]
        rupees = ",".join(pairs) + "," + last3
    return f"{sign}{rupees}.{paise}"
